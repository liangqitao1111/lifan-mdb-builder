"""理反 — 土工试验统计模块（原始数据显示版）

统计实现已与 soil_stats.py 收敛（复评 P2-5）：公共数据结构与统计工具
（IndicatorStats/StratumStats/_safe_float/_median/_statistics/_is_reasonable_value/
collect_stratum_statistics 等）统一从 soil_stats 导入，不再各自维护一套平行实现。
本模块只保留 v2 特有的 RawSample 原始样本收集 + 双工作表（原始数据/标贯原始数据）
输出逻辑；对外 API 与合并前完全兼容（actions.py 调用方不变：
v1 collect_stratum_statistics + v2 _collect_raw_samples / write_dual_sheet_workbook）。
"""

from collections import defaultdict
from dataclasses import dataclass
from typing import Dict, Optional

import openpyxl
from openpyxl.utils import get_column_letter

from config import (classify_lithology, PLASTICITY_ORDER, DENSITY_ORDER,
                    spt_to_plasticity, spt_to_density, il_to_plasticity,
                    load_project_config)

# 公共实现统一来自 soil_stats（消除 v1/v2 平行实现分叉）
from soil_stats import (
    IndicatorStats, StratumStats,                              # 数据结构（v1 同源）
    _safe_float, _median, _statistics, _is_reasonable_value,   # 统计工具（v1 同源）
    _quantile, _abs_cv,                                        # T1 标准分位 / T6 |CV| 剔除决策（v1 同源）
    _count_decimals, CV_MIN_ABS_MEAN,                          # H5 动态精度 / R1 分母保护（V2.2.4）
    _remove_outliers_iqr, _remove_outliers_by_median,          # 离群剔除（v1 同源）
    INDICATORS, INDICATOR_GROUPS, DOMAIN_RULES, REASONABLE_RANGES,
    STANDARD_VALUE_DIRECTION, CV_THRESHOLD, MAX_ITER, OUTLIER_SIGMA,
    MIN_STD_SAMPLE, PLASTICITY_RANGES,
    collect_stratum_statistics,                                # 统计入口（v1 同源，API 兼容）
)

# =============================================================================
# 旧规则备份（硬编码兜底，仅 DECIMAL_OVERRIDE 使用——v2 特有的"原始数据"表格式）
# =============================================================================

SOIL_KEYS = ['hsl','gmd','Gs','kxb','yx','sx','sxzs','yxzs','alpha','Es','phi','cohesion']

LEGACY_DECIMAL_OVERRIDE = {'hsl':1,'gmd':1,'yx':1,'sx':1,'sxzs':1,'yxzs':2,
    'phi':1,'cohesion':1,'alpha':2,'Es':2,'Gs':2,'kxb':3,'N':1,'N_corr':1}

# 小数位覆盖（v2 特有：供"原始数据/标贯原始数据"表公式与数字格式使用；
# INDICATORS/REASONABLE_RANGES 等其余指标配置已由 soil_stats 统一加载）
def _build_config():
    """重建 v2 小数位覆盖表（import 时与配置保存后各调用一次）"""
    global _cfg, _ti, DECIMAL_OVERRIDE

    _cfg = load_project_config()
    _ti = _cfg.get('试验指标', {})
    if _ti:
        DECIMAL_OVERRIDE = {k: v['小数位'] for k, v in _ti.items() if v.get('小数位') is not None}
    else:
        DECIMAL_OVERRIDE = dict(LEGACY_DECIMAL_OVERRIDE)


def reload_from_config():
    """配置保存后重建本模块派生常量（review_api._reload_config_modules 调用）"""
    _build_config()


_build_config()


# =========================================================================
# 数据结构（v2 特有：原始样本）
# =========================================================================
@dataclass
class RawSample:
    zkbh:str;qybh:str;qysd:float
    values:Dict[str,Optional[float]]
    excluded:Dict[str,bool]

# =========================================================================
# 原始样本收集 + 剔除标记
# =========================================================================
def _round_raw(v, decimals):
    """按指标动态精度舍入原始值（V2.2.4 H5；decimals=0 时保留原值）"""
    return round(v, decimals) if decimals else v


def _collect_raw_samples(da, project_type):
    """逐样本收集 + 与原版完全一致的剔除标记"""
    all_bh={b['zkbh']:b for b in da.get_all_boreholes()}
    all_strata=da.get_all_strata();all_test=da.get_all_test_full(project_type);all_spt=da.get_all_spt()

    state_count={}
    for zkbh,strata in all_strata.items():
        for s in strata:
            cb=str(s.get('tczcbh','')).strip();yb=str(s.get('tcycbh','')).strip()
            if not cb:continue
            key=f'{cb}-{yb}' if yb else cb
            st=str(s.get('tcksx','')).strip() or str(s.get('tcmsd','')).strip() or str(s.get('tcfhcd','')).strip()
            if st:state_count.setdefault(key,{});state_count[key][st]=state_count[key].get(st,0)+1
    def _inferred_state(cb,yb):
        key=f'{cb}-{yb}' if yb else cb;cnt=state_count.get(key,{})
        return max(cnt.items(),key=lambda x:x[1])[0] if cnt else ''

    raw_samples=[];groups={}
    for zkbh,strata in all_strata.items():
        bh=all_bh.get(zkbh);tests=all_test.get(zkbh,[]);spt_list=all_spt.get(zkbh,[])
        if not bh or not strata:continue
        for idx,s in enumerate(strata):
            cb=str(s.get('tczcbh','')).strip();yb=str(s.get('tcycbh','')).strip()
            if not cb:continue
            layer_label=f'{cb}-{yb}' if yb else cb
            prev_depth=strata[idx-1]['tccdsd'] if idx>0 else 0;top,bottom=prev_depth,s['tccdsd']
            layer_tests=[t for t in tests if top<t['qysd']<=bottom]
            raw_name=str(s.get('tcymc','') or s.get('tcmc','')).strip()
            if '风化' in raw_name or classify_lithology(raw_name)=='rock':continue
            name=raw_name.replace('粘','黏')
            raw_state=str(s.get('tcksx','')).strip() or str(s.get('tcmsd','')).strip() or str(s.get('tcfhcd','')).strip()
            key=(layer_label,str(s.get('tcdzsd','')).strip(),name)

            for t in layer_tests:
                values={};excluded={}
                for k in SOIL_KEYS:
                    v=t.get(k);values[k]=float(v) if v is not None and not isinstance(v,float) else v
                    excluded[k]=False
                check_state=raw_state or _inferred_state(cb,yb)
                for k in SOIL_KEYS:
                    if values[k] is not None and not _is_reasonable_value(k,values[k]):excluded[k]=True
                il=values.get('yxzs')
                if il is not None and check_state in PLASTICITY_ORDER:
                    if il_to_plasticity(il,project_type)!=check_state:
                        for k in ('yx','sx','sxzs','yxzs'):excluded[k]=True
                raw_samples.append((key,RawSample(zkbh=zkbh,qybh=str(t.get('qybh','')).strip(),
                    qysd=round(float(t.get('qysd',0) or 0),2),values=values,excluded=excluded)))
                t2={k:None if excluded[k] else values.get(k) for k in SOIL_KEYS}
                t2['_qybh'] = str(t.get('qybh','')).strip()  # 用于 CV 回填匹配
                groups.setdefault(key,[]).append(t2)

            TL=0.45
            for sp in spt_list:
                spt_top=sp['bgdsd']
                if not(spt_top>=top and spt_top+TL<=bottom):continue
                n_val=sp['bgjs'] if sp['bgjs'] and sp['bgjs']>0 else sp.get('bgxzjs',0)
                if n_val and n_val>0 and not _is_reasonable_value('N',n_val):continue
                if n_val and n_val>0 and raw_state:
                    if raw_state in PLASTICITY_ORDER:
                        if spt_to_plasticity(n_val,project_type,'不限制')!=raw_state:continue  # 浮点直判（修正击数可为小数；None=空隙跳过）
                    elif raw_state in DENSITY_ORDER:
                        st_d=spt_to_density(n_val)
                        if st_d is None or st_d!=raw_state:continue
                else:
                    # 状态缺失：按同编号最常见状态兜底
                    inferred = _inferred_state(cb, yb)
                    if inferred and inferred in PLASTICITY_ORDER:
                        if spt_to_plasticity(n_val,project_type,'不限制')!=inferred:continue  # 浮点直判（None=空隙跳过）
                    elif inferred and inferred in DENSITY_ORDER:
                        st_d=spt_to_density(n_val)
                        if st_d is None or st_d!=inferred:continue
                    else:
                        litho = classify_lithology(name)
                        if n_val and n_val>0:
                            if litho=='clay' and (n_val<2 or n_val>35):continue
                            if litho=='sand' and (n_val<3 or n_val>40):continue
                            if litho in ('fill','cavity','muck') and (n_val<1 or n_val>15):continue
                            if litho=='gravel' and (n_val<5 or n_val>80):continue
                            if litho=='rock' and (n_val<40 or n_val>200):continue
                # V2.2.2 V1：N=0（未做试验）行不进样本集——与 v1 统计及本表 _extract_spt_raw 口径一致
                if not sp.get('bgjs'):continue
                groups.setdefault(key,[]).append({'qybh':'','qysd':sp['bgdsd'],'hsl':None,'gmd':None,
                    'Gs':None,'kxb':None,'yx':None,'sx':None,'sxzs':None,'yxzs':None,
                    'alpha':None,'Es':None,'phi':None,'cohesion':None,'N':sp['bgjs'],'N_corr':sp.get('bgxzjs'),
                    '_zkbh': zkbh})  # 标贯钻孔号

    raw_by_key=defaultdict(list)
    for key,rs in raw_samples:raw_by_key[key].append(rs)
    for k in raw_by_key:raw_by_key[k].sort(key=lambda rs:(rs.zkbh,rs.qysd))

    # CV 剔除回填（与原版 _compute_group_stats 完全一致）
    for (label,era,name),samples in groups.items():
        rsk=raw_by_key.get((label,era,name),[])
        samples=[t.copy() for t in samples]
        for t in samples:
            for k in INDICATORS:
                if k=='yxzs':continue
                if t.get(k) is not None and t[k]<0:t[k]=None
        for rn,rk,rf in DOMAIN_RULES:
            if rn in name:
                for t in samples:
                    if t.get(rk) is not None and rf(t[rk]):t[rk]=None
        for gn,lk in INDICATOR_GROUPS:
            if not(set(gn)&set(SOIL_KEYS)):continue
            gn_samps=sum(1 for s in samples if any(s.get(k) is not None for k in gn))
            if gn_samps >= 6:
                # 组联动剔除（与原版 _compute_group_stats 一致）
                valid=[{**{k:s.get(k) for k in gn},'_qybh':s.get('_qybh','')} for s in samples if any(s.get(k) is not None for k in gn)]
                for kk in gn:
                    vals=[s[kk] for s in valid if s[kk] is not None];nv=len(vals)
                    if nv<4:continue
                    sv=sorted(vals);q1,q3=_quantile(sv,1),_quantile(sv,3);iqr=q3-q1
                    if iqr==0:continue
                    lo,hi=q1-1.5*iqr,q3+1.5*iqr
                    valid=[s for s in valid if s.get(kk) is None or(lo<=s[kk]<=hi)]
                # 3σ+MAD 批量剔除 + 逐个剔除 (同上)
                for _ in range(MAX_ITER):
                    if len(valid)<=2:break
                    mc,wk=0.0,None
                    for kk in gn:
                        vv=[s[kk] for s in valid if s[kk] is not None]
                        if not vv:continue
                        st=_statistics(vv)
                        c=_abs_cv(st)  # T6：剔除决策用 |CV|（负均值组不提前 break）；V2.2.4 R1 分母保护
                        if c is not None and c>mc:mc=c;wk=kk
                    if mc<=CV_THRESHOLD or wk is None:break
                    kv=[s[wk] for s in valid if s[wk] is not None];med=_median(kv)
                    if med is None or med==0:break
                    mad=_median([abs(v-med) for v in kv])
                    if mad==0:break
                    valid=[s for s in valid if s.get(wk) is None or abs(s[wk]-med)<=OUTLIER_SIGMA*1.4826*mad]
                for _ in range(MAX_ITER):
                    if len(valid)<=2:break
                    mc,wk,wk_st=0.0,None,None
                    for kk in gn:
                        vv=[s[kk] for s in valid if s[kk] is not None]
                        if not vv:continue
                        st=_statistics(vv)
                        c=_abs_cv(st)  # T6：剔除决策用 |CV|（负均值组不提前 break）
                        if c is not None and c>mc:mc=c;wk=kk;wk_st=st
                    if mc<=CV_THRESHOLD or wk is None:break
                    # V2.2.4 R1：最差指标均值趋零（|mean|<CV_MIN_ABS_MEAN）时 CV 判据
                    # 不可收敛 → 停止逐个剔除（保留样本目标，与 v1 _compute_group_stats 一致）
                    # V2.2.5 R1 残余：负均值组（含 |mean|≥0.5）同口径停止（v1 同源，
                    # mean<CV_MIN_ABS_MEAN 即含全部负均值；阶段1 极端离群剔除保留）
                    if wk_st is not None and wk_st.avg is not None and wk_st.avg<CV_MIN_ABS_MEAN:break
                    kv=[s[wk] for s in valid if s[wk] is not None];med=_median(kv)
                    if med is None:break
                    wi=max(range(len(valid)),key=lambda i:abs(valid[i].get(wk,med)-med) if valid[i].get(wk) is not None else 0)
                    valid.pop(wi)
                kept_qybhs={s['_qybh'] for s in valid}
                for rs in rsk:
                    if rs.qybh not in kept_qybhs:
                        for k in gn:
                            if rs.values.get(k) is not None:rs.excluded[k]=True
            else:
                # 个体剔除（与原版 _compute_with_direction 一致）
                for kk in gn:
                    vals=[(rs.values.get(kk),rs.qybh) for rs in rsk if rs.values.get(kk) is not None and not rs.excluded.get(kk)]
                    if len(vals)<4:continue
                    nums=[v for v,_ in vals]
                    kept_nums=set(_remove_outliers_by_median(_remove_outliers_iqr(nums)))
                    for v,qybh in vals:
                        if v not in kept_nums:
                            for rs in rsk:
                                if rs.qybh==qybh:rs.excluded[kk]=True
        for i,rs in enumerate(rsk):
            for k in INDICATORS:
                if k=='yxzs':continue
                if rs.values.get(k) is not None and rs.values[k]<0:rs.excluded[k]=True
            for rn,rk,rf in DOMAIN_RULES:
                if rn in name:
                    v=rs.values.get(rk)
                    if v is not None and rf(v):rs.excluded[rk]=True

        # ---- SPT 剔除：与原版 _compute_group_stats(['N','N_corr'], 'N') 口径一致 ----
        # 记录原版实际剔除的样本（IQR→3σ(MAD)→逐个剔除 全流程复刻），
        # 写表时不再做"距中位数最远"的事后近似（两者选中的样本集可能不一致）。
        spt_samps=[s for s in samples if any(s.get(k) is not None for k in ('N','N_corr'))]
        if len(spt_samps)>=4:
            valid=list(spt_samps)
            # 阶段0：IQR 一次清（整条剔除联动）
            for kk in ('N','N_corr'):
                vals=[s[kk] for s in valid if s.get(kk) is not None]
                nv=len(vals)
                if nv<4:continue
                sv=sorted(vals);q1,q3=_quantile(sv,1),_quantile(sv,3);iqr=q3-q1
                if iqr==0:continue
                lo,hi=q1-1.5*iqr,q3+1.5*iqr
                valid=[s for s in valid if s.get(kk) is None or (lo<=s[kk]<=hi)]
            # 阶段1：3σ(MAD) 批量剔除
            for _ in range(MAX_ITER):
                if len(valid)<=2:break
                mc,wk=0.0,None
                for kk in ('N','N_corr'):
                    vv=[s[kk] for s in valid if s.get(kk) is not None]
                    if not vv:continue
                    st=_statistics(vv)
                    c=_abs_cv(st)  # T6：剔除决策用 |CV|（负均值组不提前 break）；V2.2.4 R1 分母保护
                    if c is not None and c>mc:mc=c;wk=kk
                if mc<=CV_THRESHOLD or wk is None:break
                kv=[s[wk] for s in valid if s.get(wk) is not None];med=_median(kv)
                if med is None or med==0:break
                mad=_median([abs(v-med) for v in kv])
                if mad==0:break
                th=OUTLIER_SIGMA*1.4826*mad
                valid=[s for s in valid if s.get(wk) is None or abs(s[wk]-med)<=th]
            # 阶段2：CV 仍高时逐个剔除离中位数最远的样本（双向）
            for _ in range(MAX_ITER):
                if len(valid)<=2:break
                mc,wk,wk_st=0.0,None,None
                for kk in ('N','N_corr'):
                    vv=[s[kk] for s in valid if s.get(kk) is not None]
                    if not vv:continue
                    st=_statistics(vv)
                    c=_abs_cv(st)  # T6：剔除决策用 |CV|（负均值组不提前 break）
                    if c is not None and c>mc:mc=c;wk=kk;wk_st=st
                if mc<=CV_THRESHOLD or wk is None:break
                # V2.2.4 R1：最差指标均值趋零时 CV 判据不可收敛 → 停止逐个剔除
                # V2.2.5 R1 残余：负均值组同口径停止（v1 _compute_group_stats 同源；
                # SPT 组 N>0 恒正均值，此守卫对 SPT 无实际触发，保持四路径一致性）
                if wk_st is not None and wk_st.avg is not None and wk_st.avg<CV_MIN_ABS_MEAN:break
                kv=[s[wk] for s in valid if s.get(wk) is not None];med=_median(kv)
                if med is None:break
                worst_idx=max(range(len(valid)),
                              key=lambda i: abs(valid[i].get(wk,med)-med) if valid[i].get(wk) is not None else 0)
                valid.pop(worst_idx)
            kept={id(s) for s in valid}
            for s in spt_samps:
                if id(s) not in kept:
                    s['_spt_excluded']={'N':True,'N_corr':True}

        # V2.2.3 T2：剔除标记写回 groups——此前 marks 落在循环内的局部副本上被丢弃，
        # _extract_spt_raw 永远看不到剔除 → v2 标贯原始表与 v1 保留样本集不一致
        # （v1 已剔除的行在 v2 表仍显示为正常值）。写回后 v1 保留数 == v2 未剔除行数。
        groups[(label,era,name)] = samples

    return raw_by_key,groups

# =========================================================================
# 写入 Sheet "原始数据"
# =========================================================================
def _write_raw_data_sheet(ws,stats_list,raw_by_key):
    from openpyxl.styles import Font,Alignment,Border,Side,PatternFill
    hf=PatternFill(start_color='4472C4',end_color='4472C4',fill_type='solid')
    hfo=Font(name='宋体',size=8,color='FFFFFF',bold=True);df=Font(name='宋体',size=8)
    tf=Font(name='宋体',size=10,bold=True)
    center=Alignment(horizontal='center',vertical='center',wrap_text=True)
    thin=Border(left=Side(style='thin'),right=Side(style='thin'),top=Side(style='thin'),bottom=Side(style='thin'))

    cols=['序号','钻孔号','取样编号','深度(m)']+[INDICATORS[k][0] for k in SOIL_KEYS];n_cols=len(cols)
    for ci,h in enumerate(cols,1):
        c=ws.cell(row=1,column=ci,value=h);c.font=hfo;c.fill=hf;c.alignment=center;c.border=thin
    for ci,w in enumerate([4,16,8,8]+[7]*len(SOIL_KEYS),1):ws.column_dimensions[get_column_letter(ci)].width=w

    cur=2;row_labels=['统计个数','最大值','最小值','平均值','标准差','变异系数','标准值','中大平均值','中小平均值']
    for ss in stats_list:
        key=(ss.layer_label,ss.era_genesis,ss.name);samples=raw_by_key.get(key,[])
        if not samples:continue
        # V2.2.4 H5：原始值按指标动态精度写出（_count_decimals 上限 8，与 v1 统计一致），
        # 不再硬编码 round(v,2)——0.951→0.95 的截断使公式统计与 v1 全精度漂移
        key_decimals={}
        for k in SOIL_KEYS:
            vals=[rs.values.get(k) for rs in samples if rs.values.get(k) is not None]
            key_decimals[k]=_count_decimals(vals) if vals else DECIMAL_OVERRIDE.get(k,2)
        tc=ws.cell(row=cur,column=1,value=f'{ss.layer_label}  {ss.era_genesis}  {ss.name}  ({ss.plasticity})');tc.font=tf
        ws.merge_cells(start_row=cur,start_column=1,end_row=cur,end_column=n_cols)
        for ci in range(1,n_cols+1):ws.cell(row=cur,column=ci).border=thin
        cur+=1;stat_start=cur;ds=cur+len(row_labels);de=ds+len(samples)-1
        for ri,lb in enumerate(row_labels):
            r=cur+ri;ws.cell(row=r,column=1,value=lb).font=df
            ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=4)
            for ci in range(1,n_cols+1):ws.cell(row=r,column=ci).border=thin;ws.cell(row=r,column=ci).alignment=center
            for ki,k in enumerate(SOIL_KEYS):
                ci=5+ki;cl=get_column_letter(ci);rng=f'{cl}{ds}:{cl}{de}'
                dec=DECIMAL_OVERRIDE.get(k,2)
                fmts={0:'0',1:'0.'+'0'*dec,2:'0.'+'0'*dec,3:'0.'+'0'*dec,4:'0.000',5:'0.000',6:'0.'+'0'*dec,7:'0.'+'0'*dec,8:'0.'+'0'*dec}
                if ri==0:ws.cell(row=r,column=ci,value=f'=COUNT({rng})')
                elif ri==1:ws.cell(row=r,column=ci,value=f'=IF(COUNT({rng})=0,"",MAX({rng}))')  # V2.2.4 R2：全剔除列 MAX 输出空白（与 v1 一致）
                elif ri==2:ws.cell(row=r,column=ci,value=f'=IF(COUNT({rng})=0,"",MIN({rng}))')  # V2.2.4 R2：同上 MIN
                elif ri==3:ws.cell(row=r,column=ci,value=f'=IF(COUNT({rng})=0,"",AVERAGE({rng}))')  # T5：全剔除列 COUNT=0 不输出 #DIV/0!
                elif ri==4:ws.cell(row=r,column=ci,value=f'=IF(COUNT({rng})=0,"",STDEV({rng}))')  # V2.2.4 R2：STDEV=STDEV.S（样本标准差），兼容 Excel 2007+
                elif ri==5:ws.cell(row=r,column=ci,value=f'=IF(COUNT({rng})=0,"",IF({cl}{stat_start+3}<=0,"",ROUND({cl}{stat_start+4}/{cl}{stat_start+3},3)))')  # V2.2.4 R2：先判 COUNT 再判均值（avg 空串不再 #VALUE!）；V2.2.2 S1 负均值空白不变
                elif ri==6:
                    if STANDARD_VALUE_DIRECTION=='规范':
                        sign='+' if INDICATORS[k][1]=='low_good' else '-'
                    else:
                        sign='-'  # '模板'：与 铁三物理力学表公式.xlsx 一致，全部 1−ψ
                    p=f'(1.704/SQRT(COUNT({rng}))+4.678/COUNT({rng})^2)';c2=f'MAX(STDEV({rng})/AVERAGE({rng}),0)'
                    # V2.2.4 R2：嵌套 IF 先判 COUNT(<6→空白) 再判 AVERAGE=0，替代 OR()——
                    # OR 急切求值会把 AVERAGE(空区间) 的 #DIV/0! 传给 IF，使 COUNT<6 守卫失效
                    # V2.2.5：负均值组标准值按 ψ=0（CV 钳 0）口径——CV 项 MAX(...,0)
                    # 与 v1 _statistics 钳制一致：负均值组标准值 = 平均值，v1/v2 不再分叉
                    # （第四轮新可疑项①：v1 -0.0917 vs v2 旧公式 -0.21；与 S1 负 CV
                    # 口径同源；原始数据表与标贯原始数据表两处公式同步）。
                    ws.cell(row=r,column=ci,value=f'=IF(COUNT({rng})<6,"",IF(AVERAGE({rng})=0,"",ROUND(AVERAGE({rng})*(1{sign}{p}*{c2}),{dec})))')
                elif ri==7:ws.cell(row=r,column=ci,value=f'=IF(COUNT({rng})<6,"",ROUND((MAX({rng})+AVERAGE({rng}))/2,{dec}))')
                elif ri==8:ws.cell(row=r,column=ci,value=f'=IF(COUNT({rng})<6,"",ROUND((MIN({rng})+AVERAGE({rng}))/2,{dec}))')
                if ri in fmts:ws.cell(row=r,column=ci).number_format=fmts[ri]
                ws.cell(row=r,column=ci).font=df
        cur+=len(row_labels)
        for idx,rs in enumerate(samples,1):
            r=cur;ws.cell(row=r,column=1,value=idx).font=df;ws.cell(row=r,column=2,value=rs.zkbh).font=df
            ws.cell(row=r,column=3,value=rs.qybh).font=df;ws.cell(row=r,column=4,value=rs.qysd).font=df
            for ki,k in enumerate(SOIL_KEYS):
                ci=5+ki;v=rs.values.get(k);dec=key_decimals.get(k,2)
                if v is None:ws.cell(row=r,column=ci,value='')
                elif rs.excluded.get(k,False):ws.cell(row=r,column=ci,value=f'*{_round_raw(v,dec)}')
                else:ws.cell(row=r,column=ci,value=_round_raw(v,dec))
                ws.cell(row=r,column=ci).font=df
            for ci in range(1,n_cols+1):ws.cell(row=r,column=ci).alignment=center;ws.cell(row=r,column=ci).border=thin
            cur+=1
        cur+=1
    return ws

# =========================================================================
# 写入 Sheet "标贯原始数据"
# =========================================================================
def _extract_spt_raw(groups):
    spt_raw=defaultdict(list)
    for (label,era,name),samples in groups.items():
        for t in samples:
            n=t.get('N')
            if n and n>0:
                excl = t.get('_spt_excluded', {})
                spt_raw[(label,era,name)].append({
                    'zkbh':t.get('_zkbh',''),'depth':t.get('qysd',0),
                    'N':n,'N_corr':t.get('N_corr'),
                    'excluded_N':excl.get('N',False),'excluded_N_corr':excl.get('N_corr',False)})
    return spt_raw

def _write_spt_raw_sheet(ws,stats_list,spt_raw):
    from openpyxl.styles import Font,Alignment,Border,Side,PatternFill
    hf=PatternFill(start_color='4472C4',end_color='4472C4',fill_type='solid')
    hfo=Font(name='宋体',size=8,color='FFFFFF',bold=True);df=Font(name='宋体',size=8)
    tf=Font(name='宋体',size=10,bold=True)
    center=Alignment(horizontal='center',vertical='center',wrap_text=True)
    thin=Border(left=Side(style='thin'),right=Side(style='thin'),top=Side(style='thin'),bottom=Side(style='thin'))
    cols=['钻孔号','深度(m)','实测N','修正N'];n_cols=len(cols)
    for ci,h in enumerate(cols,1):
        c=ws.cell(row=1,column=ci,value=h);c.font=hfo;c.fill=hf;c.alignment=center;c.border=thin
    for ci in range(1,n_cols+1):ws.column_dimensions[get_column_letter(ci)].width=12
    cur=2;row_labels=['统计个数','最大值','最小值','平均值','标准差','变异系数','标准值','中大平均值','中小平均值']
    for ss in stats_list:
        key=(ss.layer_label,ss.era_genesis,ss.name);samples=spt_raw.get(key,[])
        if not samples:continue
        tc=ws.cell(row=cur,column=1,value=f'{ss.layer_label}  {ss.era_genesis}  {ss.name}  ({ss.plasticity})');tc.font=tf
        ws.merge_cells(start_row=cur,start_column=1,end_row=cur,end_column=n_cols)
        for ci in range(1,n_cols+1):ws.cell(row=cur,column=ci).border=thin
        cur+=1;stat_start=cur;ds=cur+len(row_labels);de=ds+len(samples)-1;dec=1
        for ri,lb in enumerate(row_labels):
            r=cur+ri;ws.cell(row=r,column=1,value=lb).font=df
            ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=1)
            for ci in range(1,n_cols+1):ws.cell(row=r,column=ci).border=thin;ws.cell(row=r,column=ci).alignment=center
            for ki,k in enumerate(['N','N_corr']):
                ci=3+ki;cl=get_column_letter(ci);rng=f'{cl}{ds}:{cl}{de}'
                fmts={0:'0',1:'0.'+'0'*dec,2:'0.'+'0'*dec,3:'0.'+'0'*dec,4:'0.000',5:'0.000',6:'0.'+'0'*dec,7:'0.'+'0'*dec,8:'0.'+'0'*dec}
                if ri==0:ws.cell(row=r,column=ci,value=f'=COUNT({rng})')
                elif ri==1:ws.cell(row=r,column=ci,value=f'=IF(COUNT({rng})=0,"",MAX({rng}))')  # V2.2.4 R2：全剔除列 MAX 输出空白（与 v1 一致）
                elif ri==2:ws.cell(row=r,column=ci,value=f'=IF(COUNT({rng})=0,"",MIN({rng}))')  # V2.2.4 R2：同上 MIN
                elif ri==3:ws.cell(row=r,column=ci,value=f'=IF(COUNT({rng})=0,"",AVERAGE({rng}))')  # T5：全剔除列 COUNT=0 不输出 #DIV/0!
                elif ri==4:ws.cell(row=r,column=ci,value=f'=IF(COUNT({rng})=0,"",STDEV({rng}))')  # V2.2.4 R2：STDEV=STDEV.S（样本标准差），兼容 Excel 2007+
                elif ri==5:ws.cell(row=r,column=ci,value=f'=IF(COUNT({rng})=0,"",IF({cl}{stat_start+3}<=0,"",ROUND({cl}{stat_start+4}/{cl}{stat_start+3},3)))')  # V2.2.4 R2：先判 COUNT 再判均值（avg 空串不再 #VALUE!）；V2.2.2 S1 负均值空白不变
                elif ri==6:
                    if STANDARD_VALUE_DIRECTION=='规范':
                        sign='+' if INDICATORS[k][1]=='low_good' else '-'
                    else:
                        sign='-'  # '模板'：与 铁三物理力学表公式.xlsx 一致，全部 1−ψ
                    p=f'(1.704/SQRT(COUNT({rng}))+4.678/COUNT({rng})^2)';c2=f'MAX(STDEV({rng})/AVERAGE({rng}),0)'
                    # V2.2.4 R2：嵌套 IF 先判 COUNT(<6→空白) 再判 AVERAGE=0，替代 OR()——
                    # OR 急切求值会把 AVERAGE(空区间) 的 #DIV/0! 传给 IF，使 COUNT<6 守卫失效
                    # V2.2.5：负均值组标准值按 ψ=0（CV 钳 0）口径——CV 项 MAX(...,0)
                    # 与 v1 _statistics 钳制一致：负均值组标准值 = 平均值，v1/v2 不再分叉
                    # （第四轮新可疑项①：v1 -0.0917 vs v2 旧公式 -0.21；与 S1 负 CV
                    # 口径同源；原始数据表与标贯原始数据表两处公式同步）。
                    ws.cell(row=r,column=ci,value=f'=IF(COUNT({rng})<6,"",IF(AVERAGE({rng})=0,"",ROUND(AVERAGE({rng})*(1{sign}{p}*{c2}),{dec})))')
                elif ri==7:ws.cell(row=r,column=ci,value=f'=IF(COUNT({rng})<6,"",ROUND((MAX({rng})+AVERAGE({rng}))/2,{dec}))')
                elif ri==8:ws.cell(row=r,column=ci,value=f'=IF(COUNT({rng})<6,"",ROUND((MIN({rng})+AVERAGE({rng}))/2,{dec}))')
                if ri in fmts:ws.cell(row=r,column=ci).number_format=fmts[ri]
                ws.cell(row=r,column=ci).font=df
        cur+=len(row_labels)
        for rs in samples:
            r=cur;ws.cell(row=r,column=1,value=rs.get('zkbh','')).font=df;ws.cell(row=r,column=2,value=rs['depth']).font=df
            ws.cell(row=r,column=3,value=f'*{round(rs["N"],1)}' if rs.get('excluded_N') else round(rs['N'],1) if rs['N'] else '').font=df
            ncorr=rs.get('N_corr')
            v=ncorr if ncorr else '';ws.cell(row=r,column=4,value=f'*{v}' if rs.get('excluded_N_corr') else v).font=df
            for ci in range(1,n_cols+1):ws.cell(row=r,column=ci).alignment=center;ws.cell(row=r,column=ci).border=thin
            cur+=1
        cur+=1
    return ws

# =========================================================================
# 主输出
# =========================================================================
def write_dual_sheet_workbook(stats_list, raw_by_key, groups, output_path):
    """输出双工作表 xlsx"""
    wb=openpyxl.Workbook()
    ws_raw=wb.active;ws_raw.title='原始数据';_write_raw_data_sheet(ws_raw,stats_list,raw_by_key)
    spt_raw=_extract_spt_raw(groups)
    # SPT 剔除标记已在 _collect_raw_samples 内按原版 _compute_group_stats 口径
    # 记录实际剔除样本（写表时不再做"距中位数最远"的事后近似）
    ws_spt=wb.create_sheet('标贯原始数据');_write_spt_raw_sheet(ws_spt,stats_list,spt_raw)
    wb.save(output_path)
    return output_path
