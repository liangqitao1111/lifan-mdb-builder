"""理反 — 复核规则引擎"""
import re, os
import openpyxl
from dataclasses import dataclass
from config import (
    classify_lithology, classify_soil, il_to_plasticity, spt_to_density, spt_to_plasticity,
    spt_to_weathering, dpt_to_density,
    load_project_config, normalize_state_word,
)
import config as _cfg_mod  # P2-2（Codex 复核）：DENSITY_ORDER/VERTICAL_THIN_LAYER_DENSE 按值导入会在
                           # reload_config() 重建后失效 → R-DEN-010 恒用旧值；改运行时读取 config 模块
from applog import get_logger
import dao  # V2.2.2（C4）：构造时向 DAO 注入工程类型，接线 A 类动探杆长修正

# ---- 标准地层状态参数表：列号常量（与 参数/地层标准状态参数表.xlsx 表头一一对应，
#      唯一读取器 _load_std_stratum_table 使用，消除双读取/双硬编码列号）----
_STD_TABLE_FILE = '地层标准状态参数表.xlsx'
_STD_COL_MAIN = 1      # 主层
_STD_COL_SUB = 2       # 亚层
_STD_COL_LITH = 5      # 岩土名称
_STD_COL_TCMSD = 8     # 密实度
_STD_COL_TCSID = 9     # 湿度
_STD_COL_TCKSX = 10    # 可塑性
_STD_COL_TCFHCD = 11   # 风化程度
_STD_COL_STATE = 13    # 原始状态描述


# ---- 从 TOML 读取关键词（P1-⑤：可经 reload_from_config 重建，参数中心保存后生效）----
_COLOR_WORDS = re.compile('(?:)')   # 占位，_build_keywords() 重建
_DESC_DENSITY = re.compile('(?:)')
_DESC_MOISTURE = re.compile('(?:)')
_COLOR_EXCLUDE_WORDS = []
_GENERIC_EXCLUDE_WORDS = []


def _build_keywords():
    """重建模块级关键词词表（import 时与配置保存后各调用一次）"""
    global _COLOR_WORDS, _DESC_DENSITY, _DESC_MOISTURE
    global _COLOR_EXCLUDE_WORDS, _GENERIC_EXCLUDE_WORDS
    _kw_cfg = load_project_config().get('规则_关键词', {})
    _color_words_list = _kw_cfg.get('颜色词',
        ['灰','黄','红','褐','棕','黑','白','绿','紫','青','蓝','橙',
         '粉红','暗红','肉红','砖红','紫红','黄褐','灰白','灰黄','灰绿',
         '灰褐','深灰','浅灰','棕黄','棕红','深褐',
         '黄绿','灰黑','青灰','乳白','淡黄','深黄','浅黄','暗绿','墨绿',
         '杂色','花色'])
    _COLOR_WORDS = re.compile(f'({chr(124).join(_color_words_list)})')
    _desc_density_words = _kw_cfg.get('密实度词',
        ['密实度','密实','中密','稍密','松散','稍密状','中密状','密实状','松散状'])
    _DESC_DENSITY = re.compile(f'({chr(124).join(_desc_density_words)})')
    _desc_moisture_words = _kw_cfg.get('湿度词',
        ['湿度','稍湿','很湿','饱和','干燥','潮湿'])
    _DESC_MOISTURE = re.compile(f'({chr(124).join(_desc_moisture_words)})')
    _COLOR_EXCLUDE_WORDS = _kw_cfg.get('颜色排除词',
        ['灰岩', '灰质', '石灰', '泥灰'])
    _GENERIC_EXCLUDE_WORDS = _kw_cfg.get('排除词',
        ['未做试验', '未做', '未测定', '未测', '无', '不', '未'])


def reload_from_config():
    """配置保存后重建本模块派生的关键词常量（review_api._reload_config_modules 调用）"""
    _build_keywords()


_build_keywords()

# ---- V2.2.2（C8）关键词排除词（最小化修复，保留模糊包含容错特色）----
# 模糊包含匹配是项目容错特色（"灰岩质黏土"含"灰"、描述含"密实度"三字即放行）；
# 但否定语境下会误放行（"该层密实度未做试验"被当密实度标注、颜色缺失被"灰岩"字样
# 放行）。新增排除词表：命中关键词与排除词重叠/紧邻（中间仅允许空格/顿号）时不算匹配。
# 颜色词仅用岩性类排除词（灰岩/灰质/石灰/泥灰等），避免"灰黄色，无杂质"等
# 正常颜色描述被通用否定词误伤；密实度/湿度词用试验缺失类否定词。
# V2.2.4（E3）：收紧"紧邻"语义，修复"中密未见地下水/稍密未见地下水"被误判为
# 否定语境（R-DEN-005/006 H 级误报）——多字否定短语（未做/未测/未做试验/未测定）
# 直接相邻才生效；单字否定词（无/不/未）仅当关键词为字段概念词（密实度/湿度）、
# 或否定词为"不"（"不密实"直接否定状态）、或"未见X"前置（"未见密实度"）时生效。
# "中密未见地下水"中"未"修饰的是地下水而非密实度，不再触发否定。
# 字段概念词：否定词修饰的是"字段本身"（"密实度未见记录/未见密实度"= 无密实度资料）；
# 状态词（中密/稍密/密实/松散/稍湿/饱和 等）被"未/无"邻接多为误伤
# （"中密未见地下水"= 中密 + 未见地下水，密实度信息仍成立）。
_DENSITY_CONCEPT_WORDS = frozenset(['密实度'])
_MOISTURE_CONCEPT_WORDS = frozenset(['湿度'])
# E3：相邻间隔仅允许 空格/全角空格/顿号；其余字符（含逗号"中密，未见地下水"）一律视为不相邻
_GAP_SEP_RE = re.compile(r'^[ \u3000、]*$')
# V2.2.5（N2）：前置"未见X"否定句式的 '见' 与关键词之间容忍空格/全角空格
# （'未见 密实度' 与 '未见密实度' 同为否定；此前仅认紧贴的 '见'）
_LEFT_SEE_RE = re.compile(r'^见[ \u3000]*$')


def _kw_excluded(desc, start, end, excludes, concept_words=(), matched=''):
    """关键词匹配段 (start,end) 是否被排除词否定

    E3 收紧后的否定判定（满足其一即否定）：
      1) 重叠：排除词与关键词段重叠（灰岩 包住 灰）；
      2) 多字否定短语（未做试验/未做/未测定/未测）与关键词直接相邻
         （中间仅允许空格/顿号，如"密实度 未做试验"/"密实度、未做试验"）；
      3) 单字否定词（无/不/未）相邻时，仅当关键词为字段概念词（密实度/湿度，
         如"密实度未见记录"）、否定词为"不"（"不密实"直接否定状态）、
         或前置"未见X"（"未见密实度"）时生效。
    "中密未见地下水"中的"未"修饰地下水，不满足上述任一条件 → 不再误判为否定语境。
    """
    for ex in excludes:
        if not ex:
            continue
        single = len(ex) == 1
        pos = desc.find(ex)
        while pos != -1:
            ex_start, ex_end = pos, pos + len(ex)
            if ex_start < end and ex_end > start:      # 重叠：灰岩 包住 灰
                return True
            # 直接相邻，或中间仅空格/顿号（E3：不再接受 ≤2 任意无标点字符的"短间隔"）
            right_adj = ex_start == end or (ex_start > end and _GAP_SEP_RE.match(desc[end:ex_start]))
            left_gap = desc[ex_end:start] if ex_end < start else ''
            # 单字否定词前置间隔额外允许 '见'（"未见X"否定句式：未见密实度/未见中密；
            # V2.2.5（N2）容忍 '见' 与关键词间的空格：'未见 密实度' 同样识别）
            left_adj = ex_end == start or (ex_end < start and (_GAP_SEP_RE.match(left_gap) or _LEFT_SEE_RE.match(left_gap)))
            if right_adj or left_adj:
                if not single:                          # 多字否定短语：未做试验/未做/未测…
                    return True
                # 单字否定词：仅字段概念词 / "不"直接否定状态 / "未见X"前置 时生效
                if matched in concept_words or ex == '不' or _LEFT_SEE_RE.match(left_gap):
                    return True
            pos = desc.find(ex, pos + 1)
    return False


def _kw_search(desc, pattern, excludes, concept_words=()):
    """关键词匹配；被排除词否定的匹配段不计，全部被否定返回 None"""
    if not desc:
        return None
    for m in pattern.finditer(desc):
        if _kw_excluded(desc, m.start(), m.end(), excludes, concept_words, m.group(0)):
            continue
        return m
    return None


def _has_color_in_desc(desc):
    """检查描述文本中是否包含颜色关键词（灰岩/石灰等岩性词除外）"""
    return bool(_kw_search(desc, _COLOR_WORDS, _COLOR_EXCLUDE_WORDS))


def _desc_has_density(desc):
    """检查描述文本中是否包含密实度相关内容（"密实度未做试验"等否定语境除外）"""
    return bool(_kw_search(desc, _DESC_DENSITY, _GENERIC_EXCLUDE_WORDS, _DENSITY_CONCEPT_WORDS))


def _desc_has_moisture(desc):
    """检查描述文本中是否包含湿度相关内容（"湿度未做试验"等否定语境除外）"""
    return bool(_kw_search(desc, _DESC_MOISTURE, _GENERIC_EXCLUDE_WORDS, _MOISTURE_CONCEPT_WORDS))


def _norm_state_key(v):
    """地层编号键归一化（加载器 _load_param_state_table 与查询端共用）

    V2.2.2（C5）：'01' → '1'、'1.0' → '1'、数值 1.0 → '1'（整数化字符串表示
    统一为 str(int(float(v)))）；非整数/非数值（如 '1.5'、'A'）保留原串。
    消除前导零/浮点地层编号在 R-PLS-001/002 查表时的静默漏报。
    """
    s = str(v).strip()
    try:
        f = float(s)
    except (ValueError, TypeError):
        return s
    return str(int(f)) if f == int(f) else s


@dataclass
class ReviewIssue:
    """复核问题数据类"""
    rule_id: str
    risk_level: str
    layer_index: int
    field: str
    message: str
    layer_label: str = ''
    ref_value: float = 0.0   # 辅助匹配值（标贯/动探深度等）


@dataclass
class RuleDef:
    """规则定义（用于集中注册表，方便审查/过滤/扩展）"""
    rule_id: str
    risk_level: str
    category: str       # 'check' | 'density' | 'cross' | 'plasticity'
    description: str
    check_fn: callable  # 检查函数，接收 ctx dict 参数


def _layer_label(tczcbh: str, tcycbh: str) -> str:
    if not tczcbh: return ''
    if tcycbh: return f"{tczcbh}-{tcycbh}"
    return str(tczcbh)


class RuleEngine:

    def __init__(self, project_type: str = 'A', max_plasticity: str = '硬塑',
                 use_std_stratum: bool = False):
        # V3.0.3（Agent1-2）注释：max_plasticity 默认"硬塑"封顶为【项目设定】——
        # A 类最大塑性上限=硬塑（残积土等场景封顶，N 判出"坚硬"的黏土层按硬塑处理），
        # 非规范条款；B 类 4 档（无可塑）同样封顶硬塑（TOML 已标注"项目特色，勿改"），
        # 勿按规范表"修正"回坚硬（详见 config.spt_to_plasticity 注释）。
        self.project_type = project_type
        self.max_plasticity = max_plasticity
        self.use_std_stratum = use_std_stratum
        # V2.2.2（C4）：DAO 层按工程类型决定是否做 A 类动探杆长修正；
        # 既有调用点（get_dpt_data/get_all_dpt）不传 project_type，此处注入模块级值
        dao.set_project_type(project_type)
        # 从 TOML 读取塑性顺序（A类5级/B类4级）
        _cfg = load_project_config()
        _order_key = 'A类' if project_type == 'A' else 'B类'
        _plastic_list = (_cfg.get(_order_key, {})
                         .get('塑性状态', {})
                         .get('塑性顺序',
                              ['流塑','软塑','可塑','硬塑','坚硬'] if project_type == 'A'
                              else ['流塑','软塑','硬塑','坚硬']))
        self._order = _plastic_list
        self._max_idx = self._order.index(max_plasticity) if max_plasticity in self._order else 99
        self._std_stratum = self._load_standard_stratum()  # {(tczcbh, tcycbh): {field: value}}
        self._rules = self._build_rule_registry()

    def get_rules(self):
        """公开访问：返回当前启用的规则注册表副本（避免外部直接读私有 _rules）。

        由 Agent A（外壳/UI 域）按协调约定新增；本方法不改动任何现有逻辑。
        """
        return list(self._rules)

    def _build_rule_registry(self):
        """构建规则注册表（读取 TOML 规则配置，支持 启用/禁用/等级覆盖）"""
        base = [
            # ---- 检查类 ----
            ('R-CHK-001', 'H', 'check', '颜色字段为空', self._chk_tcys),
            ('R-CHK-002', 'H', 'check', '岩石地层风化程度为空', self._chk_tcfhcd),
            ('R-CHK-003', 'H', 'check', '描述含超出本层范围的深度', self._chk_desc_depth),
            # ---- 密实度/可塑性/风化程度一致性 ----
            ('R-DEN-001', 'H', 'density', '标贯N值→密实度不符', self._den_spt_match),
            ('R-DEN-002', 'H', 'density', '动探N63.5→密实度不符', self._den_dpt_match),
            ('R-DEN-003', 'L', 'density', '有密实度标注但该层无标贯/动探', self._den_no_test),
            ('R-DEN-004', 'H', 'density', '黏性土不应标注密实度', self._den_clay_density),
            ('R-DEN-005', 'H', 'density', '砂土缺少密实度', self._den_sand_missing),
            ('R-DEN-006', 'H', 'density', '碎石土缺少密实度', self._den_gravel_missing),
            ('R-DEN-009', 'H', 'density', '砂土/碎石土/填土缺少湿度描述', self._den_moisture_missing),
            ('R-DEN-007', 'H', 'density', '标贯N→可塑性不符', self._den_spt_plasticity),
            ('R-DEN-008', 'H', 'density', '标贯N→风化程度不符', self._den_spt_weathering),
            # ---- 交叉检查 ----
            ('R-CRS-001', 'H', 'cross', '淤泥不应标注密实度', self._crs_muck),
            ('R-CRS-010', 'M', 'cross', '盐渍土应注明含盐量', self._crs_saline),
            ('R-CRS-011', 'M', 'cross', '冻土应注明含冰量', self._crs_frozen),
            ('R-CRS-012', 'M', 'cross', '标准地层N值不达标', self._crs_std_stratum),
            # ---- 土工试验判别 ----
            # V2.2.5（R-PLS）：001/002 语义不同（TCKSX vs IL），拆分独立函数——
            # 原两规则绑定同一函数使每条问题重复上报 2 次；拆分后各规则只产出自己的
            # rule_id，单侧启用时禁用方不输出。
            ('R-PLS-001', 'H', 'plasticity', '参数表塑性状态与层位不符', self._pls_state_mismatch),
            ('R-PLS-002', 'M', 'plasticity', 'IL试验结果与层位状态不符', self._pls_il_mismatch),
            ('R-CRS-002', 'M', 'plasticity', '固结试验孔隙比异常偏高', self._crs_compressibility_void),
            ('R-CRS-003', 'M', 'plasticity', '固结试验孔隙比反常(极高)', self._crs_compressibility_mismatch),
            # ---- 颗分试验判别 ----
            ('R-GRS-001', 'M', 'cross', '颗分定名与地层岩土名称不符', self._grs_grain_size_match),
            # ---- 纵向顺序检查（真实规则函数，注册表统一调度；TOML 启停/等级覆盖在下方绑定） ----
            ('R-DEN-010', 'M', 'density', '深部密实度比上层更松', None),  # fn 在下文按生效等级绑定
        ]

        # 读取 TOML 规则配置
        _cfg = load_project_config()
        _rules_cfg = _cfg.get('规则', {})
        _stat_cfg = _cfg.get('试验指标_统计', {})
        _global_cfg = _cfg.get('规则', {}).get('_全局', {})

        # 获取全局默认启用设置：_默认启用=true 时，未列出的规则默认启用；false 时默认禁用
        _default_enabled = True
        _default_entry = _rules_cfg.get('_默认启用', None)
        if isinstance(_default_entry, dict):
            _default_enabled = _default_entry.get('启用', True)
        elif isinstance(_default_entry, bool):
            _default_enabled = _default_entry

        # 缓存阈值供规则方法使用（实例属性命名，非类常量风格）
        self._rule_threshold = {
            '深度公差': _global_cfg.get('描述深度公差', 0.1),
            '孔隙比异常': _stat_cfg.get('孔隙比异常', 1.1),
            '孔隙比极高': _stat_cfg.get('孔隙比极高', 2.5),
            # V2.2.3（S5）：0.95 次级提示阈值入 TOML（新增键 孔隙比_次级阈值）
            '孔隙比次级': _stat_cfg.get('孔隙比_次级阈值', 0.95),
        }

        result = []
        for code, level, category, desc, fn in base:
            # TOML 配置覆盖（id 用 code 直接匹配）
            rc = _rules_cfg.get(code, {})
            if rc.get('启用', _default_enabled) is False:
                continue  # 禁用则跳过
            eff_level = rc.get('等级', level)
            eff_category = rc.get('类别', category)
            eff_desc = rc.get('说明', desc)
            if code == 'R-DEN-010':
                # 纵向检查：把 TOML 生效等级绑定进真实规则函数，
                # 注册表统一调度（含启用/禁用），不再 noop 占位 + 反向查找等级
                fn = (lambda ctx, lv=eff_level: self._den_vertical_check(ctx, lv))
            result.append(RuleDef(code, eff_level, eff_category, eff_desc, fn))
        return result

    @staticmethod
    def _safe_layer(layer: dict) -> dict:
        def s(v, d=''): return str(v) if v is not None else d
        def f(v, d=0.0):
            if v is None: return d
            try: return float(v)
            except (ValueError, TypeError): return d
        def i(v, d=0):
            if v is None: return d
            try: return int(v)
            except (ValueError, TypeError): return d
        return {
            'tczcbh': s(layer.get('tczcbh')), 'tcycbh': s(layer.get('tcycbh')),
            'tcxh': i(layer.get('tcxh')), 'tccdsd': f(layer.get('tccdsd')),
            'tchd': f(layer.get('tchd')), 'tcymc': s(layer.get('tcymc')),
            'tcmc': s(layer.get('tcmc')), 'tcys': s(layer.get('tcys')),
            # V2.2.5（S4）：tcksx 带状写法归一化（'硬塑状' → '硬塑'），与 tcmsd 同口径
            'tcksx': normalize_state_word(layer.get('tcksx')),
            # V2.2.3（S7）：tcmsd 带状写法归一化（'稍密状' → '稍密'），
            # 与 DAO _row_to_stratum 同口径，覆盖非 DAO 直接构造层位数据的路径
            'tcmsd': normalize_state_word(layer.get('tcmsd')),
            'tcfhcd': s(layer.get('tcfhcd')), 'tcsid': s(layer.get('tcsid')),
            'tcms': s(layer.get('tcms')),
        }

    @staticmethod
    def _safe_spt(spt: dict):
        def f(v, d=0.0):
            if v is None: return d
            try: return float(v)
            except (ValueError, TypeError): return d
        def i(v, d=0):
            if v is None: return d
            try: return int(v)
            except (ValueError, TypeError): return d
        return {'bgdsd': f(spt.get('bgdsd')), 'bggc': f(spt.get('bggc')),
                'bgjs': i(spt.get('bgjs')), 'bgxzjs': f(spt.get('bgxzjs'))}

    def review_strata(self, strata, spt_data=None, dpt_data=None, test_data=None):
        """对一组地层执行全套复核，返回问题列表"""
        if spt_data is None: spt_data = []
        if dpt_data is None: dpt_data = []
        if test_data is None: test_data = []
        strata = [self._safe_layer(l) for l in strata]
        # 确保按层底深度排序（防御：调用者可能传入未排序数据）
        strata.sort(key=lambda x: float(x.get('tccdsd', 0)))
        spt_data = [self._safe_spt(s) for s in spt_data]
        if dpt_data:
            dpt_data = [{**d, 'dtdsd': float(d.get('dtdsd', 0) or 0),
                         'dtjs': int(d.get('dtjs', 0) or 0),
                         # V2.2.2（C6）：修正击数浮点直判（与 C1 同口径），不再 int() 截断
                         'dtxzjs': float(d.get('dtxzjs', 0) or 0)} for d in dpt_data]
        if test_data:
            test_data = [{**t, 'qysd': float(t.get('qysd', 0) or 0),
                          'yxzs': float(t.get('yxzs', 0)) if t.get('yxzs') is not None else None} for t in test_data]
        issues = []
        # V24：跨层标贯异常检查——标贯深度穿越层范围时优先报异常（在逐层击数规则前执行）
        issues.extend(self._check_spt_cross_layer(strata, spt_data))
        for i, layer in enumerate(strata):
            depth = layer['tccdsd']
            prev_depth = strata[i - 1]['tccdsd'] if i > 0 else 0
            layer_spt = [s for s in spt_data if prev_depth < s['bgdsd'] <= depth]
            layer_dpt = [d for d in dpt_data if prev_depth < d['dtdsd'] <= depth]
            layer_test = [t for t in test_data if prev_depth < t['qysd'] <= depth]
            ctx = {'layer': layer, 'index': i, 'prev_depth': prev_depth,
                   'spt': layer_spt, 'dpt': layer_dpt, 'test': layer_test,
                   'all_strata': strata, 'all_spt': spt_data}

            for rule in self._rules:
                try:
                    result = rule.check_fn(ctx)
                except Exception:
                    # 单规则异常不中断整孔复核（v43）：记录日志并跳过该规则，
                    # 避免异常数据导致全库复核 500
                    get_logger().exception('规则 %s 执行异常（层号 %s，层索引 %s），已跳过',
                                           rule.rule_id, layer.get('tczcbh', ''), i)
                    continue
                if result:
                    if isinstance(result, list):
                        issues.extend(result)
                    else:
                        issues.append(result)
        # R-DEN-010 纵向检查已注册为真实规则，随注册表统一调度（无需手动接线）
        return issues

    def _make_issue(self, rid, level, ctx, field, msg):
        layer = ctx['layer']
        return ReviewIssue(rule_id=rid, risk_level=level, layer_index=ctx['index'],
                           field=field, message=msg,
                           layer_label=_layer_label(layer.get('tczcbh', ''), layer.get('tcycbh', '')))

    def _check_spt_cross_layer(self, strata, spt_data):
        """V24 跨层标贯异常：标贯记录深度应落在某一地层的 (层顶, 层底] 区间内；
        不在任何层范围（层间空隙 / 超出最浅~最深范围）→ 判定异常。
        优先级高于标贯击数规则（R-DEN-001/007/008）：在逐层规则前执行。"""
        issues = []
        if not strata or not spt_data:
            return issues
        ranges = []
        prev = 0.0
        for layer in strata:
            depth = float(layer.get('tccdsd', 0) or 0)
            ranges.append((prev, depth))
            prev = depth
        lo_all = ranges[0][0] if ranges else 0
        hi_all = ranges[-1][1] if ranges else 0
        for s in spt_data:
            d = float(s.get('bgdsd', 0) or 0)
            if d <= 0:
                continue
            if not any(lo < d <= hi for lo, hi in ranges):
                issues.append(ReviewIssue(
                    rule_id='R-SPT-001', risk_level='H', layer_index=-1,
                    field='BGDSD',
                    message=f'标贯深度 {_norm_state_key(d)}m 不在任何地层深度范围内（层序 {_norm_state_key(lo_all)}~{_norm_state_key(hi_all)}m），疑似跨层或深度误录',
                    ref_value=d))
        return issues

    # ---- 颜色/风化程度漏填检查 ----
    def _chk_tcys(self, ctx):
        l = ctx['layer']
        if classify_lithology(l['tcymc']) == 'cavity': return
        if not l['tcys']:
            # 描述列中存在颜色信息则不判风险
            desc = l.get('tcms', '')
            if _has_color_in_desc(desc):
                return
            return self._make_issue('R-CHK-001', 'H', ctx, 'TCYS', '颜色字段为空')

    def _chk_tcfhcd(self, ctx):
        l = ctx['layer']
        if classify_lithology(l['tcymc']) == 'cavity': return
        if classify_lithology(l['tcymc']) == 'rock' and not l['tcfhcd']:
            # V2.2.3（S4）：名称本身含风化/残积信息（全风化岩/强风化岩/中风化岩/
            # 微风化岩/残积土等）时，风化程度已由名称表达，字段为空不再误报 H；
            # 仅名称不含风化信息且字段为空才提示补充。
            if any(kw in l['tcymc'] for kw in ('风化', '残积')):
                return
            return self._make_issue('R-CHK-002', 'H', ctx, 'TCFHCD', '岩石地层风化程度为空')

    # ---- 描述内容深度范围检查 ----
        # V3.0.3（Agent1-15）注释：R-CHK-003 为【经验值（项目特色）】检查——
        # 描述栏深度数字超出本层范围（公差「描述深度公差」默认 0.1m）疑为误录，
        # 属数据卫生规则，非规范条文；阈值由 TOML「规则._全局.描述深度公差」配置。
    def _chk_desc_depth(self, ctx):
        """检查描述中的深度数值是否超出本层深度范围"""
        l = ctx['layer']
        desc = l.get('tcms', '')
        if not desc:
            return

        prev_depth = ctx.get('prev_depth', 0)
        layer_bottom = l['tccdsd']
        # P2（dao）：层底深度为 0（源库该字段为 NULL，dao 转 0）时范围 [0,0] 会让
        # 描述里任何深度数字都误报 R-CHK-003——无有效深度范围，跳过该层检查
        if layer_bottom <= 0:
            return

        # 提取所有"数字+m/米"或"数字~数字+m/米"的深度值
        # 使用 finditer 获取位置，过滤掉 cm/mm 的误匹配
        raw_pattern = r'(\d+\.?\d*)\s*~\s*(\d+\.?\d*)\s*[m米]|(\d+\.?\d*)\s*[m米]'
        out_of_range = []

        for m in re.finditer(raw_pattern, desc):
            full_match = m.group()
            match_end = m.end()

            # 定位单位字符(m/米)在原描述中的位置，跳过末尾空白
            unit_pos = match_end - 1
            while unit_pos > m.start() and desc[unit_pos] in (' ', '\t'):
                unit_pos -= 1
            unit_char = desc[unit_pos] if unit_pos >= 0 else ''

            # 检查单位前的字符（跳过空白），排除 cm / mm
            before_pos = unit_pos - 1
            while before_pos >= 0 and desc[before_pos] in (' ', '\t'):
                before_pos -= 1
            before_char = desc[before_pos].lower() if before_pos >= 0 else ''

            # 检查单位后的字符，排除 mm (第二个m)
            after_pos = unit_pos + 1
            after_char = desc[after_pos].lower() if after_pos < len(desc) else ''

            # 排除: "cm", "mm" (单位前是c或m), "mm" (单位后紧跟m)
            if before_char in ('c', 'm') or after_char == 'm':
                continue

            # 排除纯数字后无单位的情况（如 thickness=5 不含 m）
            if unit_char not in ('m', '米'):
                continue

            # 提取数值
            groups = m.groups()
            if groups[2]:  # 单数值: match group 3
                val = float(groups[2])
                tol = self._rule_threshold.get('深度公差', 0.1)
                if val < prev_depth - tol or val > layer_bottom + tol:
                    out_of_range.append(f'{val:.1f}m')
            elif groups[1]:  # 范围: match group 1, 2
                v1, v2 = float(groups[0]), float(groups[1])
                tol = self._rule_threshold.get('深度公差', 0.1)
                for v in (v1, v2):
                    if v < prev_depth - tol or v > layer_bottom + tol:
                        label = f'{v:.1f}m'
                        if label not in out_of_range:
                            out_of_range.append(label)

        if out_of_range:
            items = '、'.join(str(x) for x in out_of_range[:5])
            return self._make_issue('R-CHK-003', 'H', ctx, 'TCMS',
                f'描述含超出本层范围({prev_depth:.1f}~{layer_bottom:.1f}m)的深度: {items}')

    # ---- 密实度一致性检查（逐条标贯/动探独立判定） ----
    def _den_spt_match(self, ctx):
        l = ctx['layer']
        if self.use_std_stratum:  # 勾选标准地层表 → 跳过N值推导
            return
        if classify_lithology(l['tcymc']) not in ('sand', 'gravel') or not ctx['spt']:
            return
        if not l['tcmsd']:
            return
        results = []
        for spt in ctx['spt']:
            bgjs = spt.get('bgjs')
            bgxzjs = spt.get('bgxzjs')
            # V3.0.1 业务口径修正（GB 50021-2001 表3.3.9）：砂土/碎石土密实度判别用
            # 【标准贯入实测击数】，不做杆长修正；仅当实测值缺失时回退修正击数。
            # 注：修正值仍用于承载力查表/可塑性/风化等（勿混用）。
            n = bgjs if bgjs is not None and bgjs > 0 else (bgxzjs or 0)
            if n <= 0:  # 空标贯（未录入），不参与密实度对比（与 R-DEN-007 同防护，避免除零/误报"松散"）
                continue
            expected = spt_to_density(n)
            if expected is None:  # 区间外（配置异常）不判定
                continue
            if l['tcmsd'] != expected:
                issue = self._make_issue('R-DEN-001', 'H', ctx, 'BGJS',
                    f'N={n}对应{expected}，当前密实度为{l["tcmsd"]}（按实测击数）')
                issue.ref_value = spt.get('bgdsd', 0)
                results.append(issue)
        return results if results else None

        # V3.0.3（Agent1-15）注释：R-DEN-002 适用 GB50021 表3.3.8-1 碎石土密实度分类——
        # 该表适用于【平均粒径≤50mm 且最大粒径<100mm】的碎石土（表注；平均粒径>50mm 或
        # 最大粒径>100mm 应用超重型 N120 表3.3.8-2）。本项目数据无粒径字段，按适用处理，
        # 属数据卫生经验口径，不做岩性细分。
    def _den_dpt_match(self, ctx):
        l = ctx['layer']
        if classify_lithology(l['tcymc']) != 'gravel' or not ctx['dpt']:
            return
        if not l['tcmsd']:
            return
        results = []
        for dpt in ctx['dpt']:
            # V3.0.2（B4）：仅重型动探（DTLX='2'/中文'重型'）参与 N63.5 密实度判别
            # （GB50021 表B.0.1 仅适用重型 N63.5；轻型 N10/超重型 N120 不参与）。
            # 未标记 is_heavy 的调用方数据按重型处理（兼容手写 dict / 既有测试）。
            if not dpt.get('is_heavy', True):
                continue
            n = dpt.get('dtxzjs')
            # V3.0.2（A5/D）：无修正值时不再回退实测 N63.5——实测值查"修正后"
            # 口径表属口径错配（密实度系统性偏高估一档）；修正值由 dao 按工程
            # 类型（A/B 同表B.0.1）自动计算，库中仍无修正值则跳过判别。
            if n is None or n <= 0:  # 空动探/无修正值，不参与对比（同 R-DEN-001 防护）
                continue
            expected = dpt_to_density(n)
            if expected is None:  # 区间外（配置异常）不判定
                continue
            if l['tcmsd'] != expected:
                issue = self._make_issue('R-DEN-002', 'H', ctx, 'TCMSD',
                    f'N63.5={n}对应{expected}，当前为{l["tcmsd"]}')
                issue.ref_value = dpt.get('dtdsd', 0)
                results.append(issue)
        return results if results else None

    def _den_no_test(self, ctx):
        l = ctx['layer']; lt = classify_lithology(l['tcymc'])
        # V2.2.3（S8）：按层判断（ctx['spt']/ctx['dpt'] 已是本层数据）——
        # 该层有密实度标注但无标贯/动探数据支撑即报，不再依赖全孔 all_spt
        # （原实现整孔无任何标贯时静默不报，与 TOML 说明"该层无标贯或动探数据
        # 支撑该判定"不符）；全库噪音由 TOML 默认禁用开关控制。
        if lt in ('sand', 'gravel') and l['tcmsd'] and not ctx['spt'] and not ctx['dpt']:
            return self._make_issue('R-DEN-003', 'L', ctx, 'BGJS',
                f'{l["tcymc"]}有密实度标注但该层无标贯/动探数据支撑')

    def _den_clay_density(self, ctx):
        l = ctx['layer']
        # V2.2.2（C7）：muck（淤泥/软土）由 R-CRS-001"淤泥不应标注密实度"专项覆盖，
        # 此处只判 clay，消除同一层同一字段（TCMSD）两条 H 级规则重复双报（无信息增量）。
        if classify_lithology(l['tcymc']) == 'clay' and l['tcmsd']:
            return self._make_issue('R-DEN-004', 'H', ctx, 'TCMSD', f'{l["tcymc"]}应用可塑性描述，非密实度')

    def _den_sand_missing(self, ctx):
        l = ctx['layer']
        if classify_lithology(l['tcymc']) == 'sand' and not l['tcmsd']:
            # 描述列中已有密实度相关内容则不判风险
            if _desc_has_density(l.get('tcms', '')):
                return
            return self._make_issue('R-DEN-005', 'H', ctx, 'TCMSD', '砂土缺少密实度')

    def _den_gravel_missing(self, ctx):
        l = ctx['layer']
        if classify_lithology(l['tcymc']) == 'gravel' and not l['tcmsd']:
            # 描述列中已有密实度相关内容则不判风险
            if _desc_has_density(l.get('tcms', '')):
                return
            return self._make_issue('R-DEN-006', 'H', ctx, 'TCMSD', '碎石土缺少密实度')

    def _den_moisture_missing(self, ctx):
        """检查砂土/碎石土/填土湿度缺失，描述列中有湿度相关内容则不判风险"""
        l = ctx['layer']
        if classify_lithology(l['tcymc']) not in ('sand', 'gravel', 'fill'): return
        if not l['tcsid']:
            if _desc_has_moisture(l.get('tcms', '')):
                return
            return self._make_issue('R-DEN-009', 'H', ctx, 'TCSID', '砂土/碎石土/填土缺少湿度描述')

    # ---- 黏性土标贯 → 可塑性检查 ----
    def _den_spt_plasticity(self, ctx):
        if self.use_std_stratum:  # 勾选标准地层表 → 跳过N值推导
            return
        l = ctx['layer']
        is_clay = classify_lithology(l['tcymc']) in ('clay', 'muck')
        is_residual = l.get('tcfhcd', '') == '残积土'
        if not (is_clay or is_residual) or not ctx['spt']:
            return
        if not l['tcksx']:
            return
        results = []
        for spt in ctx['spt']:
            bgjs = spt.get('bgjs')
            bgxzjs = spt.get('bgxzjs')
            # V3.0.2 业务口径修正：黏性土可塑性按【实测标贯击数 N】判定——广东省土层判定
            # 指标注"N 为未经杆长修正的标准贯入锤击数"，GB 50021-2001 表3.3.11 黏性土
            # 状态按液性指数 IL 判定，N 判可塑性属地区经验表；与 soil_stats 统计口径
            # （实测优先）及 R-DEN-001（V3.0.1）同源；仅实测缺失时回退修正击数。
            n = bgjs if bgjs is not None and bgjs > 0 else (bgxzjs or 0)
            if n <= 0:  # 空标贯（未录入），不参与可塑性对比
                continue
            expected = spt_to_plasticity(n, self.project_type, self.max_plasticity)
            if expected is None:  # 区间空隙（修正击数浮点）→ 不判定
                continue
            if l['tcksx'] != expected:
                issue = self._make_issue('R-DEN-007', 'H', ctx, 'TCKSX',
                    f'N={n}对应{expected}，当前可塑性为{l["tcksx"]}')
                issue.ref_value = spt.get('bgdsd', 0)
                results.append(issue)
        return results if results else None

    def _den_spt_weathering(self, ctx):
        if self.use_std_stratum:  # 勾选标准地层表 → 跳过N值推导
            return
        l = ctx['layer']
        if not ctx['spt']:
            return
        if not l['tcfhcd']:
            return
        if l['tcfhcd'] not in ('全风化', '强风化'):
            return
        results = []
        for spt in ctx['spt']:
            bgjs = spt.get('bgjs')
            bgxzjs = spt.get('bgxzjs')
            # V3.0.2 业务口径修正：风化程度按【实测标贯击数 N】判定——DBJ 15-31-2016
            # 第4章"以实测标贯击数 N' 划分花岗岩类岩石风化程度"；广东省土层判定指标 4.1
            # 通篇 N 均为未经杆长修正；与 R-DEN-001/007（V3.0.1/V3.0.2）同源；
            # 仅实测缺失时回退修正击数。
            n = bgjs if bgjs is not None and bgjs > 0 else (bgxzjs or 0)
            if n <= 0:  # 空标贯（未录入），不参与风化程度对比（同 R-DEN-001 防护）
                continue
            expected = spt_to_weathering(n, self.project_type)
            if expected is None:  # 区间空隙 → 不判定
                continue
            if l['tcfhcd'] != expected:
                issue = self._make_issue('R-DEN-008', 'H', ctx, 'TCFHCD',
                    f'N={n}对应{expected}，当前风化为{l["tcfhcd"]}')
                issue.ref_value = spt.get('bgdsd', 0)
                results.append(issue)
        return results if results else None

    # ---- 交叉检查 ----
    def _crs_muck(self, ctx):
        l = ctx['layer']
        if classify_lithology(l['tcymc']) == 'muck' and l['tcmsd']:
            return self._make_issue('R-CRS-001', 'H', ctx, 'TCMSD', '淤泥不应标注密实度')

    def _crs_saline(self, ctx):
        if self.project_type != 'B': return
        if classify_lithology(ctx['layer']['tcymc']) == 'saline':
            return self._make_issue('R-CRS-010', 'M', ctx, 'TCYMC', '盐渍土应注明含盐量和盐渍化程度')

    def _crs_frozen(self, ctx):
        if self.project_type != 'B': return
        if classify_lithology(ctx['layer']['tcymc']) == 'frozen':
            return self._make_issue('R-CRS-011', 'M', ctx, 'TCYMC', '冻土应注明含冰量级别')

    # ---- 标准地层一致性检查 ----
    _STD_TABLE_CACHE = None  # 唯一读取器 _load_std_stratum_table 的结果缓存（进程级共享）

    def _load_std_stratum_table(self):
        """唯一的标准地层表读取器（消除双读取/双硬编码列号，复评 P2-6）

        读取 参数/地层标准状态参数表.xlsx，返回
        {(主层, 亚层): {'tcmsd','tcsid','tcksx','tcfhcd','lith','state'}}，
        键为原始字符串（消费方各自做自己的键归一化）。
        列号统一来自模块级 _STD_COL_* 常量，模板列序变化只改一处。
        """
        if self._STD_TABLE_CACHE is not None:
            return self._STD_TABLE_CACHE
        base = os.path.dirname(os.path.dirname(os.path.dirname(__file__)))
        path = os.path.join(base, '参数', _STD_TABLE_FILE)
        if not os.path.exists(path):
            self._STD_TABLE_CACHE = {}
            return {}
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
            ws = wb.active
            table = {}
            for r in range(2, ws.max_row + 1):
                main_v = ws.cell(row=r, column=_STD_COL_MAIN).value
                if main_v is None:
                    continue
                main_s = str(main_v).strip()
                sub_s = str(ws.cell(row=r, column=_STD_COL_SUB).value or '').strip()
                def _cell(col):
                    v = ws.cell(row=r, column=col).value
                    return str(v).strip() if v is not None else ''
                table[(main_s, sub_s)] = {
                    'lith': _cell(_STD_COL_LITH),
                    'tcmsd': _cell(_STD_COL_TCMSD),
                    'tcsid': _cell(_STD_COL_TCSID),
                    'tcksx': _cell(_STD_COL_TCKSX),
                    'tcfhcd': _cell(_STD_COL_TCFHCD),
                    'state': _cell(_STD_COL_STATE),
                }
            wb.close()
            self._STD_TABLE_CACHE = table
            return table
        except Exception:
            get_logger().exception('_load_std_stratum_table 读取失败: %s', path)
            self._STD_TABLE_CACHE = {}
            return {}

    def _load_standard_stratum(self):
        """加载统一地层标准状态参数表，返回 {(主层, 亚层): {密实度, 湿度, 可塑性, 风化程度}}

        列 8-11（密实度/湿度/可塑性/风化程度）来自唯一读取器 _load_std_stratum_table。
        """
        table = self._load_std_stratum_table()
        lookup = {}
        for (main, sub), row in table.items():
            if not main or not sub:
                continue
            std = {}
            for field in ('tcmsd', 'tcsid', 'tcksx', 'tcfhcd'):
                vs = row.get(field, '')
                if vs and vs != '-':
                    std[field] = vs
            if std:
                lookup[(main, sub)] = std
        return lookup

    def _crs_std_stratum(self, ctx):
        """勾选模式：按标准地层表判定标贯N值是否达标

        流程：地层编号 → 标准地层.xls → 标准状态 → 标贯N值对比
        """
        if not self.use_std_stratum:
            return
        if not self._std_stratum or not ctx['spt']:
            return
        l = ctx['layer']
        key = (l['tczcbh'], l['tcycbh'])
        std = self._std_stratum.get(key)
        if not std:
            return

        lt = classify_lithology(l['tcymc'])
        results = []

        for spt in ctx['spt']:
            bgjs = spt.get('bgjs')
            bgxzjs = spt.get('bgxzjs')
            # V3.0.2 业务口径修正：标准地层对照的塑性/风化/密实度统一用【实测击数】
            # （广东表 N 未经杆长修正；DBJ 15-31 以实测 N' 划分风化；GB 50021 表3.3.9
            # 密实度用实测 N）——与 R-DEN-001/007/008 完全同口径，不再分叉；
            # 仅实测缺失时回退修正击数。
            n = bgjs if bgjs is not None and bgjs > 0 else (bgxzjs or 0)
            if n <= 0:
                continue

            # 标准表塑性对照（黏性土）
            std_pls = std.get('tcksx', '')
            if (lt in ('clay', 'muck') or l.get('tcfhcd', '') == '残积土') and std_pls:
                n_state = spt_to_plasticity(n, self.project_type, self.max_plasticity)
                if n_state is None or n_state != std_pls:
                    issue = self._make_issue('R-CRS-012', 'M', ctx, 'BGJS',
                        f'N={n}→{n_state}，标准地层{key[0]}-{key[1]}应为"{std_pls}"')
                    issue.ref_value = spt.get('bgdsd', 0)
                    results.append(issue)

            # 标准表密实度对照（砂土/碎石土）——V3.0.1 起用实测击数（GB 50021 表3.3.9）；
            # V3.0.2 循环顶部 n 已统一为实测击数，n_dens 直接复用 n（口径同 R-DEN-001）
            std_dens = std.get('tcmsd', '')
            if lt in ('sand', 'gravel') and std_dens:
                n_dens = n
                n_state = spt_to_density(n_dens)
                if n_state is not None and n_state != std_dens:
                    # 区间空隙（None）不判定，但不得跳过本条标贯的风化对照（V39）
                    issue = self._make_issue('R-CRS-012', 'M', ctx, 'BGJS',
                        f'N={n_dens}→{n_state}，标准地层{key[0]}-{key[1]}应为"{std_dens}"')
                    issue.ref_value = spt.get('bgdsd', 0)
                    results.append(issue)

            # 标准表风化程度对照（全/强风化）
            std_wea = std.get('tcfhcd', '')
            if l.get('tcfhcd', '') in ('全风化', '强风化') and std_wea:
                n_state = spt_to_weathering(n, self.project_type)
                if n_state is None or n_state != std_wea:
                    issue = self._make_issue('R-CRS-012', 'M', ctx, 'BGJS',
                        f'N={n}→{n_state}，标准地层{key[0]}-{key[1]}应为"{std_wea}"')
                    issue.ref_value = spt.get('bgdsd', 0)
                    results.append(issue)

        return results if results else None

    # ---- 参数表状态查询（供土工试验判别使用） ----
    def _load_param_state_table(self):
        """加载统一地层标准状态参数表，返回 {(编号, 子编号): (状态, 岩性), ...}

        列5(岩土名称)+列13(原始状态描述) 来自唯一读取器 _load_std_stratum_table；
        键做 int 归一化（'0'/'01' → '0'/'1'），与历史口径一致。
        """
        table = self._load_std_stratum_table()
        result = {}
        for (main_s, sub_s), row in table.items():
            # V2.2.2（C5）：与查询端共用 _norm_state_key（'01'/'1.0' → '1'）
            b_key = _norm_state_key(main_s)
            sub_key = _norm_state_key(sub_s)
            s_val = row.get('state', '')
            if s_val:
                # 处理组合状态，保证切出完整状态词
                first = s_val.split('、')[0].strip()
                if '-' in first:
                    candidate = first.split('-')[0].strip()
                    known = {'流塑', '软塑', '可塑', '硬塑', '坚硬',
                             '密实', '中密', '稍密', '松散'}
                    if candidate not in known:
                        candidate = first
                    s_val = candidate
                else:
                    s_val = first
            result[(b_key, sub_key)] = (s_val, row.get('lith', ''))
        return result

    # ---- 土工试验判别：液塑限试验 ----
    # V2.2.5（R-PLS）：原 001/002 两规则绑定同一函数 _pls_liquidity_test，注册表把
    # 同一函数调度两次 → 每条问题重复上报 2 次（且单侧启用时仍输出另一侧）。
    # 拆分语义：001 查"层位 TCKSX vs 参数表标准"，002 查"试验IL推导状态 vs 层位 TCKSX"；
    # 共享前置 _pls_std_state（判据、参数表查询），各规则只产出自己的 rule_id。
    def _pls_std_state(self, ctx):
        """R-PLS 共享前置：适用性判据 + 参数表标准状态查询；不适用返回 None"""
        l = ctx['layer']
        if not ctx['test']:
            return None
        is_clay = classify_lithology(l['tcymc']) in ('clay', 'muck')
        if not is_clay:
            return None
        cb = _norm_state_key(l.get('tczcbh', ''))
        yb = _norm_state_key(l.get('tcycbh', ''))
        if not cb:
            return None
        param_tbl = self._load_param_state_table()
        std_state, _std_lith = param_tbl.get((cb, yb), ('', ''))
        if not std_state:
            return None
        return cb, yb, std_state

    def _pls_state_mismatch(self, ctx):
        """R-PLS-001：层位可塑性标注(tcksx)与参数表标准不一致 → 风险在TCKSX"""
        pre = self._pls_std_state(ctx)
        if pre is None:
            return
        cb, yb, std_state = pre
        l = ctx['layer']
        actual_tcksx = l.get('tcksx', '')   # 钻孔实填状态（_safe_layer 已归一化带状写法）
        results = []
        for t in ctx['test']:
            il = t.get('yxzs')
            if il is None:
                continue
            il_state = il_to_plasticity(il, self.project_type)  # 由IL推导的状态
            if il_state is None:
                continue  # 区间空隙（配置可编辑）→ 不判定（Codex 复核 P1-1 补防）
            if not actual_tcksx:
                continue  # 缺数据则不评判
            if actual_tcksx != std_state:
                issue = self._make_issue('R-PLS-001', 'H', ctx, 'TCKSX',
                    f'参数表({cb}-{yb})标准="{std_state}"，层位可塑性标注="{actual_tcksx}"，试验IL={il:.2f}→{il_state}')
                issue.ref_value = t.get('qybh', '')
                results.append(issue)
        return results if results else None

    def _pls_il_mismatch(self, ctx):
        """R-PLS-002：层位可塑性与参数表一致，但试验IL推导状态与标注不符 → 风险在IL"""
        pre = self._pls_std_state(ctx)
        if pre is None:
            return
        cb, yb, std_state = pre
        l = ctx['layer']
        actual_tcksx = l.get('tcksx', '')   # 钻孔实填状态（_safe_layer 已归一化带状写法）
        results = []
        for t in ctx['test']:
            il = t.get('yxzs')
            if il is None:
                continue
            il_state = il_to_plasticity(il, self.project_type)  # 由IL推导的状态
            if il_state is None:
                continue  # 区间空隙（配置可编辑）→ 不判定（Codex 复核 P1-1 补防）
            if not actual_tcksx:
                continue  # 缺数据则不评判
            if actual_tcksx == std_state and il_state != actual_tcksx:
                issue = self._make_issue('R-PLS-002', 'M', ctx, 'YXZS',
                    f'层位可塑性标注="{actual_tcksx}"（符合参数表），但试验IL={il:.2f}对应{il_state}')
                issue.ref_value = t.get('qybh', '')
                results.append(issue)
        return results if results else None

    # ---- 土工试验判别：固结试验 ----
    def _crs_compressibility_void(self, ctx):
        # V3.0.3（Agent1-14）注释：e₀ 阈值（1.1 异常 / 0.95 次级）为【经验值（项目特色）】——
        # 规范压缩性按压缩系数 a1-2 判定（高压缩 a1-2≥0.5MPa⁻¹），e₀ 仅为代理近似，无直接规范来源；
        # 1.1 与 0.95 区间重叠属刻意设计（e₀∈(0.95,1.1] 报次级提示，>1.1 报主提示），保留并标注。
        """R-CRS-002：根据固结试验孔隙比 e₀ 判别压缩性异常"""
        l = ctx['layer']
        if not ctx['test']:
            return
        is_clay = classify_lithology(l['tcymc']) in ('clay', 'muck')
        if not is_clay:
            return
        results = []
        for t in ctx['test']:
            e0 = t.get('kxb')
            if e0 is None:
                continue
            # e₀ > 阈值 → 高压缩性黏土，需验证
            thresh = self._rule_threshold.get('孔隙比异常', 1.1)
            if e0 > thresh:
                issue = self._make_issue('R-CRS-002', 'M', ctx, 'TCYMC',
                    f'固结试验e₀={e0:.3f}，属高压缩性，应核实与层位岩性名称{ l["tcymc"]}的一致性')
                issue.ref_value = t.get('qybh', '')
                results.append(issue)
            elif e0 > self._rule_threshold.get('孔隙比次级', 0.95):
                issue = self._make_issue('R-CRS-002', 'M', ctx, 'TCYMC',
                    f'固结试验e₀={e0:.3f}，偏高，建议结合压缩系数判断压缩性')
                issue.ref_value = t.get('qybh', '')
                results.append(issue)
        return results if results else None

    def _crs_compressibility_mismatch(self, ctx):
        # V3.0.3（Agent1-14）注释：e₀>2.5（估算 a1-2≈0.2×e₀）为【经验值（项目特色）】，
        # 非规范条款（规范压缩性按 a1-2 判定）；保留既有行为，仅补充依据说明。
        """R-CRS-003：检查固结试验孔隙比与层位岩性名称是否反常"""
        l = ctx['layer']
        if not ctx['test']:
            return
        is_clay = classify_lithology(l['tcymc']) in ('clay', 'muck')
        if not is_clay:
            return
        results = []
        for t in ctx['test']:
            e0 = t.get('kxb')
            if e0 is None:
                continue
            # 用孔隙比估算压缩系数 a1-2 ≈ 0.2 × e₀
            est_a = round(0.2 * e0, 2)
            # 高压缩性：a1-2 ≥ 0.5 MPa⁻¹ → e₀ ≥ 2.5（但阈值比较按 TOML 说明
            # "e₀ > 2.5" 严格大于执行，e₀=2.5 恰不报，与 R-CRS-002 同口径）
            # 中压缩性：0.1 ≤ a1-2 < 0.5 → 0.5 ≤ e₀ < 2.5
            thresh = self._rule_threshold.get('孔隙比极高', 2.5)
            if e0 > thresh:
                issue = self._make_issue('R-CRS-003', 'M', ctx, 'TCYMC',
                    f'固结试验e₀={e0:.3f}异常偏高，估算a1-2≈{est_a}，需复核原始数据')
                issue.ref_value = t.get('qybh', '')
                results.append(issue)
        return results if results else None

    # ---- 颗分试验判别 ----
    def _grs_grain_size_match(self, ctx):
        """R-GRS-001：按颗分数据定名，与地层岩土名称对比

        仅勾选"启用土工试验判别"时执行。
        需要该层有颗分试验数据（r0075字段非空）。
        """
        l = ctx['layer']
        if not ctx['test']:
            return
        # 归一化 粘/黏：与 soil_stats 的 replace('粘','黏') 口径一致，避免定名比对假阳性
        layer_name = str(l.get('tcymc', '')).replace('粘', '黏')
        if not layer_name:
            return

        results = []
        for t in ctx['test']:
            # 检查是否有颗分数据
            if t.get('r0075') is None:
                continue
            sp = (t.get('r2_05') or 0) + (t.get('r05_025') or 0) + (t.get('r025_0075') or 0)
            cp = t.get('r0075') or 0
            ip = t.get('sxzs')
            wl = t.get('yx')
            r2_05 = t.get('r2_05')
            r05_025 = t.get('r05_025')
            r025_0075 = t.get('r025_0075')
            r20_2 = t.get('r20_2')

            soil_name = classify_soil(sp, 0, cp, ip, wl, self.project_type,
                                       r2_05=r2_05, r05_025=r05_025,
                                       r025_0075=r025_0075, r20_2=r20_2,
                                       r0075=t.get('r0075'))
            if soil_name and soil_name != layer_name:
                issue = self._make_issue('R-GRS-001', 'M', ctx, 'TCYMC',
                    f'颗分定名="{soil_name}"，层位名称="{layer_name}"')
                issue.ref_value = t.get('qybh', '')
                results.append(issue)
        return results if results else None

    # ---- 纵向顺序检查（R-DEN-010 真实规则，注册表统一调度）----
        # V3.0.3（Agent1-15）注释：R-DEN-010 为【经验值（项目特色）】检查——
        # "深部砂土/碎石土密实度不应比上层更松"为数据卫生经验规则（砂土密实度总体随深度
        # 增大而提高，无直接规范条文）；最小厚度阈值由 TOML「公用.纵向检查.砂土纵向检查最小厚度」配置。
    def _den_vertical_check(self, ctx, level):
        """R-DEN-010：深部砂土/碎石土密实度比上层更松（厚度超过阈值时判定）

        作为真实规则注册进 _build_rule_registry，随各层 ctx 逐层调度：
        对第 i 层（i>0）与上一层比较；第 0 层无上层直接跳过。
        等级由注册表绑定 TOML 覆盖值（不再 noop 占位 + 反向查找）。
        """
        i = ctx['index']
        all_strata = ctx.get('all_strata', [])
        if i <= 0 or i >= len(all_strata):
            return
        prev, curr = all_strata[i - 1], all_strata[i]
        if (classify_lithology(prev['tcymc']) in ('sand', 'gravel')
                and classify_lithology(curr['tcymc']) in ('sand', 'gravel')
                and prev['tcmsd'] and curr['tcmsd']):
            po = _cfg_mod.DENSITY_ORDER.get(prev['tcmsd'], 0)
            co = _cfg_mod.DENSITY_ORDER.get(curr['tcmsd'], 0)
            if co < po and curr['tchd'] > _cfg_mod.VERTICAL_THIN_LAYER_DENSE:
                return self._make_issue('R-DEN-010', level, ctx, 'TCMSD',
                    f'深部密实度{curr["tcmsd"]}比上层{prev["tcmsd"]}更松')
        return None
