"""理反 — 全库复核问题导出（xlsx）

把规则引擎产出的 ReviewIssue 列表导出为 Excel 工作表，列：
钻孔编号 / 规则ID / 等级 / 层号 / 问题描述 / 位置 / 建议。

- 仅依赖 openpyxl，不依赖 GUI/数据库；
- 支持按等级过滤（H / M / L）；
- “建议”列优先取 rule_desc_map（{rule_id: 描述}），缺省给出通用提示。

用法（代码内调用）：
    from issue_exporter import export_issues_to_xlsx
    n = export_issues_to_xlsx(issues_by_zkbh, path, level_filter='H', rule_desc_map=...)
"""
import os

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADERS = ['钻孔编号', '规则ID', '等级', '层号', '问题描述', '位置', '建议']
LEVEL_LABELS = {'H': '高', 'M': '中', 'L': '低'}
_WIDTHS = [22, 10, 8, 10, 60, 26, 34]

_HEADER_FONT = Font(name='宋体', size=10, bold=True, color='FFFFFF')
_HEADER_FILL = PatternFill('solid', start_color='4472C4', end_color='4472C4')
_THIN = Side(style='thin', color='D9D9D9')
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_LEVEL_FILL = {
    'H': PatternFill('solid', start_color='FCEBEB', end_color='FCEBEB'),
    'M': PatternFill('solid', start_color='FAEEDA', end_color='FAEEDA'),
    'L': PatternFill('solid', start_color='F3F4F6', end_color='F3F4F6'),
}


def _position_text(issue):
    """由 issue 生成“位置”文本：优先 层号+深度，其次参照值/字段名"""
    parts = []
    layer_label = getattr(issue, 'layer_label', '') or ''
    if layer_label:
        parts.append(f'层 {layer_label}')
    ref = getattr(issue, 'ref_value', None)
    if ref not in (None, 0, '', 0.0):
        if isinstance(ref, (int, float)) and not isinstance(ref, bool):
            parts.append(f'深度 {float(ref):.2f}m')
        else:
            parts.append(f'参照 {ref}')
    field = getattr(issue, 'field', '') or ''
    if field:
        parts.append(field)
    return '，'.join(parts)


def _suggestion_text(issue, rule_desc_map=None):
    """生成“建议”文本：规则描述优先，缺省通用提示"""
    rule_id = getattr(issue, 'rule_id', '') or ''
    if rule_desc_map:
        desc = rule_desc_map.get(rule_id, '')
        if desc:
            return desc
    return f'请按规则 {rule_id} 复核该记录'


def natural_key(zkbh):
    """钻孔编号自然排序（ZK2 排在 ZK10 前）。

    全仓唯一实现：actions.py（钻孔列表/CAD 对比）与导出共用，
    收敛 P3-新-8 “排序 key 四处复制”问题。
    """
    import re
    return [int(p) if p.isdigit() else p for p in re.split(r'(\d+)', str(zkbh))]


def export_issues_to_xlsx(issues_by_zkbh, path, level_filter=None, rule_desc_map=None):
    """导出 {zkbh: [ReviewIssue, ...]} 为 xlsx。

    参数：
        issues_by_zkbh: dict，钻孔编号 → ReviewIssue 对象列表
        path: 输出文件路径（支持中文路径）
        level_filter: None/'H'/'M'/'L'，None 表示导出全部等级
        rule_desc_map: 可选 {rule_id: 规则说明}，用于“建议”列
    返回：实际写入的问题行数（不含表头）
    """
    if level_filter:
        level_filter = str(level_filter).strip().upper()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '复核问题'

    for c, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = _BORDER
    for i, w in enumerate(_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    rows_written = 0
    for zkbh in sorted(issues_by_zkbh.keys(), key=natural_key):
        for iss in issues_by_zkbh[zkbh]:
            level = (getattr(iss, 'risk_level', '') or '').strip().upper()
            if level_filter and level != level_filter:
                continue
            ws.append([
                zkbh,
                getattr(iss, 'rule_id', '') or '',
                level,
                getattr(iss, 'layer_label', '') or '',
                getattr(iss, 'message', '') or '',
                _position_text(iss),
                _suggestion_text(iss, rule_desc_map),
            ])
            r = ws.max_row
            for c in range(1, len(HEADERS) + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = _BORDER
                cell.alignment = Alignment(vertical='center', wrap_text=(c >= 5))
            if level in _LEVEL_FILL:
                ws.cell(row=r, column=3).fill = _LEVEL_FILL[level]
            rows_written += 1

    if rows_written:
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = f'A1:{get_column_letter(len(HEADERS))}{rows_written + 1}'
    wb.save(path)
    return rows_written


if __name__ == '__main__':
    # 无 GUI 自测：构造假 issue 数据导出到临时目录
    import sys
    import tempfile

    from dataclasses import dataclass

    @dataclass
    class _FakeIssue:
        rule_id: str
        risk_level: str
        layer_label: str
        message: str
        field: str = ''
        ref_value: float = 0.0

    fake = {
        'ZK-1': [
            _FakeIssue('R-DEN-001', 'H', '1-2', '标贯N值→密实度不符', 'TCKSX', 12.5),
            _FakeIssue('R-CHK-001', 'M', '2', '颜色字段为空', 'TCYS', 0),
        ],
        'ZK-2': [
            _FakeIssue('R-GRS-001', 'L', '3-1', '颗分定名与地层岩土名称不符', 'TCYMC', 0),
        ],
    }
    out = os.path.join(tempfile.gettempdir(), 'issue_exporter_selftest.xlsx')
    n = export_issues_to_xlsx(fake, out, rule_desc_map={'R-DEN-001': '检查标贯修正'})
    wb2 = openpyxl.load_workbook(out)
    ws2 = wb2.active
    print(f'自测：写入 {n} 行，工作簿行数 {ws2.max_row}（应含表头共 {n + 1} 行）')
    assert n == 3 and ws2.max_row == 4, '自测失败'
    n2 = export_issues_to_xlsx(fake, out, level_filter='H')
    print(f'自测：按 H 过滤写入 {n2} 行（应为 1）')
    assert n2 == 1
    os.remove(out)
    print('issue_exporter 自测通过')
    sys.exit(0)
