"""理反 — 基本承载力计算模块 (bearing_capacity.py)

根据岩土类型从铁建承载力计算表查表/插值计算查表值（列S），
从广清永标准地层参数表匹配建议值（列T）。

铁建承载力计算表来源：TB 10012-2001 铁路工程地质勘察规范
"""

import os
import re
from applog import get_logger
from config import load_project_config

_log = get_logger()

# ---- 承载力计算配置（从 TOML 加载） ----
_bc_cfg = load_project_config().get('承载力计算', {})
# 黄土液限分段阈值：新黄土 WL 分段查表用
_loess_wl_thresholds = _bc_cfg.get('黄土液限分段', [28, 32])
# 老黄土 e 分段阈值：老黄土 e 分段查表用
_loess_old_e_thresholds = _bc_cfg.get('老黄土e分段', [0.7, 0.8, 0.9])


# ============================================================
# 1D/2D 线性插值
# ============================================================

def _is_nan(v):
    """NaN 显式检测（None/非数值按 False 处理，仅 float('nan') 命中）"""
    return isinstance(v, float) and v != v


def _interp_1d(x, xs, ys):
    """一维线性插值。x 入参，xs/ys 断点-值对。超出范围取边界值。

    V2.2.3（S10）：NaN 输入不再被 min/max 比较语义静默 clamp 到边界——
    检测到 NaN 返回 None 并记日志（调用方按"无法计算"处理为空值），
    正常数值路径行为不变。
    """
    if not xs or len(xs) != len(ys):
        return 0
    if _is_nan(x):
        _log.warning('_interp_1d 收到 NaN 入参，返回 None（xs=%r）', xs)
        return None
    if x <= xs[0]:
        return ys[0]
    if x >= xs[-1]:
        return ys[-1]
    for i in range(len(xs) - 1):
        if xs[i] <= x < xs[i + 1]:
            dx = xs[i + 1] - xs[i]
            return ys[i] - (x - xs[i]) * (ys[i] - ys[i + 1]) / dx if dx > 0 else ys[i]
    return ys[-1]


def _interp_2d(row_val, col_val, row_keys, col_keys, matrix):
    """二维双线性插值。

    row_val: 行指标值（如 e）, col_val: 列指标值（如 IL）
    row_keys: 行断点列表, col_keys: 列断点列表
    matrix: matrix[ri][ci] 对应 row_keys[ri] × col_keys[ci] 的值
    """
    if not row_keys or not col_keys or not matrix:
        return 0
    # V2.2.3（S10）：NaN 输入不再被 clamp 到边界——显式检测返回 None 并记日志
    if _is_nan(row_val) or _is_nan(col_val):
        _log.warning('_interp_2d 收到 NaN 入参，返回 None（row_val=%r col_val=%r）', row_val, col_val)
        return None

    # 行 clamp
    rv = max(row_keys[0], min(row_keys[-1], row_val))
    # 列 clamp
    cv = max(col_keys[0], min(col_keys[-1], col_val))

    # 找行区间
    ri = 0
    for i in range(len(row_keys) - 1):
        if row_keys[i] <= rv <= row_keys[i + 1]:
            ri = i
            break

    # 找列区间
    ci = 0
    for i in range(len(col_keys) - 1):
        if col_keys[i] <= cv <= col_keys[i + 1]:
            ci = i
            break

    x1, x2 = row_keys[ri], row_keys[ri + 1]
    y1, y2 = col_keys[ci], col_keys[ci + 1]

    q11 = matrix[ri][ci]
    q12 = matrix[ri][ci + 1]
    q21 = matrix[ri + 1][ci]
    q22 = matrix[ri + 1][ci + 1]

    tx = (rv - x1) / (x2 - x1) if x2 > x1 else 0
    ty = (cv - y1) / (y2 - y1) if y2 > y1 else 0

    return (q11 * (1 - tx) * (1 - ty) +
            q21 * tx * (1 - ty) +
            q12 * (1 - tx) * ty +
            q22 * tx * ty)


# ============================================================
# 查表数据 — 实时读取铁建承载力计算表.xlsx（单源，无硬编码）
# ============================================================

def _load_lookup_from_excel():
    """实时读取 铁建承载力计算表.xlsx，返回所有查表数据

    每次调取重新打开 Excel，修改后即时生效，无需重启。
    返回 None 表示文件不存在或读取失败。

    返回结构:
      {'q4clay': {'e': [...], 'il': [...], 'il_max': {...}, 'matrix': [[...]]},
       'silt':    {...},
       'soft':    {'w': [...], 'sigma': [...]},
       'q3clay':  {'es': [...], 'sigma': [...]},
       'residual':{'es': [...], 'sigma': [...]},
       'loess':   {'e': [...], 'w24': [...], 'w28': [...], 'w32': [...],
                   'm24': [...], 'm28': [...], 'm32': [...]},
       'loess_old': {'matrix': [[...]]}}
    """
    import openpyxl
    base = os.path.dirname(os.path.dirname(os.path.dirname(__file__)))
    path = os.path.join(base, '参数', '铁建承载力计算表.xlsx')
    if not os.path.exists(path):
        return None

    wb = openpyxl.load_workbook(path, data_only=True)
    tables = {}

    # 辅助：从一行提取间隔列（B, D, F, ...）的有效数值
    def _extract_alt_cols(ws, row, start_col, max_col=None):
        vals = []
        max_col = max_col or ws.max_column
        for c in range(start_col, max_col + 1, 2):
            v = ws.cell(row=row, column=c).value
            if v is not None and isinstance(v, (int, float)):
                vals.append(v)
            else:
                vals.append(None)
        return vals

    # --- Q4黏性土 ---
    ws = wb['Q4黏性土']
    il_vals = _extract_alt_cols(ws, 8, 2)  # IL headers at row 8, cols B/D/F/...
    il_vals = [v for v in il_vals if v is not None]
    e_rows = [9, 10, 11, 12, 13, 14, 15]
    e_vals = []
    matrix = []
    il_max_map = {}
    for r in e_rows:
        e_v = ws.cell(row=r, column=1).value
        if e_v is None:
            continue
        e_vals.append(e_v)
        row_data = []
        last_il = 0
        for i, c in enumerate(range(2, 2 + len(il_vals) * 2, 2)):
            v = ws.cell(row=r, column=c).value
            if v is not None and isinstance(v, (int, float)):
                row_data.append(v)
                last_il = il_vals[i] if i < len(il_vals) else last_il
            else:
                row_data.append(None)
        matrix.append(row_data)
        il_max_map[e_v] = last_il
    tables['q4clay'] = {'e': e_vals, 'il': il_vals, 'il_max': il_max_map, 'matrix': matrix}

    # --- 粉土 ---
    ws = wb['粉土']
    w_vals = _extract_alt_cols(ws, 8, 2)
    w_vals = [v for v in w_vals if v is not None]
    silt_rows = [9, 10, 11, 12, 13, 14]
    silt_e = []
    silt_matrix = []
    silt_w_max = {}
    for r in silt_rows:
        e_v = ws.cell(row=r, column=1).value
        if e_v is None:
            continue
        silt_e.append(e_v)
        row_data = []
        last_w = 0
        for i, c in enumerate(range(2, 2 + len(w_vals) * 2, 2)):
            v = ws.cell(row=r, column=c).value
            if v is not None and isinstance(v, (int, float)):
                row_data.append(v)
                last_w = w_vals[i] if i < len(w_vals) else last_w
            else:
                row_data.append(None)
        silt_matrix.append(row_data)
        silt_w_max[e_v] = last_w
    tables['silt'] = {'e': silt_e, 'w': w_vals, 'w_max': silt_w_max, 'matrix': silt_matrix}

    # --- 软土 (1D) ---
    ws = wb['软土']
    soft_w = []
    soft_sigma = []
    for c in range(2, 10):
        wv = ws.cell(row=6, column=c).value
        sv = ws.cell(row=7, column=c).value
        if wv is not None and sv is not None and isinstance(wv, (int, float)) and isinstance(sv, (int, float)):
            soft_w.append(wv)
            soft_sigma.append(sv)
    tables['soft'] = {'w': soft_w, 'sigma': soft_sigma}

    # --- Q3以前黏性土 (1D) ---
    ws = wb['Q3以前黏性土']
    q3_es = []
    q3_sigma = []
    for c in range(2, 10):
        ev = ws.cell(row=6, column=c).value
        sv = ws.cell(row=7, column=c).value
        if ev is not None and sv is not None and isinstance(ev, (int, float)) and isinstance(sv, (int, float)):
            q3_es.append(ev)
            q3_sigma.append(sv)
    tables['q3clay'] = {'es': q3_es, 'sigma': q3_sigma}

    # --- 残积黏性土 (1D) ---
    ws = wb['残积黏性土']
    res_es = []
    res_sigma = []
    for c in range(2, 12):
        ev = ws.cell(row=6, column=c).value
        sv = ws.cell(row=7, column=c).value
        if ev is not None and sv is not None and isinstance(ev, (int, float)) and isinstance(sv, (int, float)):
            res_es.append(ev)
            res_sigma.append(sv)
    tables['residual'] = {'es': res_es, 'sigma': res_sigma}

    # --- 新黄土 (WL=24/28/32, e×W) ---
    ws = wb['新黄土']
    # P1-9 修复：W 表头统一读第 10 行（C10..O10 = 5..35）。
    # 原实现 WL=32 块从第 29 行（首行 σ0 数值）读列键，导致 W 永远无法落入列区间、
    # 结果恒取 W=5 档（如 WL=32/e=0.9/W=20 误得 260 而非 200）。
    w_vals = _extract_alt_cols(ws, 10, 3)
    w_vals = [v for v in w_vals if v is not None]
    loess_e = []
    for r in [11, 12, 13, 14]:
        ev = ws.cell(row=r, column=2).value  # e in col B
        if ev is not None:
            loess_e.append(ev)

    def _read_loess_block(data_rows):
        """读一个 WL 分块，按数据有效列截断（WL=24/28 块数据只到 W=30，
        原实现给空列补 0 造成 W∈(30,35] 被插值到 0 的"幻影 0 列"）。"""
        block = []
        ncols = len(w_vals)
        for r in data_rows:
            row = []
            last = 0
            for i, c in enumerate(range(3, 3 + len(w_vals) * 2, 2)):
                v = ws.cell(row=r, column=c).value
                row.append(v if isinstance(v, (int, float)) else None)
                if row[-1] is not None:
                    last = i + 1
            block.append(row)
            ncols = min(ncols, last)
        ncols = max(ncols, 1)
        return [row[:ncols] for row in block], ncols

    m24, n24 = _read_loess_block([11, 12, 13, 14])   # WL=24: rows 11-14
    m28, n28 = _read_loess_block([20, 21, 22, 23])   # WL=28: rows 20-23
    m32, n32 = _read_loess_block([29, 30, 31, 32])   # WL=32: rows 29-32（共享 W 表头）
    tables['loess'] = {
        'e': loess_e,
        'w24': w_vals[:n24], 'w28': w_vals[:n28], 'w32': w_vals[:n32],
        'm24': m24, 'm28': m28, 'm32': m32,
    }

    # --- 老黄土 (离散查表) ---
    ws = wb['老黄土']
    old_matrix = []
    for r in [10, 11, 12]:
        row = []
        for c in range(2, 6):
            v = ws.cell(row=r, column=c).value
            row.append(v if isinstance(v, (int, float)) else 0)
        old_matrix.append(row)
    tables['loess_old'] = {'matrix': old_matrix}

    wb.close()
    return tables


# ============================================================
# 岩土类型 → 查表函数映射
# ============================================================

def _ragged_2d(row_val, col_val, row_keys, col_keys, matrix, row_col_max):
    """不规则矩阵二维插值（逐行限界，不进行外推）

    row_col_max: {row_key: max_valid_col_value}
    对每个行，先限制 col_val 不超过该行的最大有效列值，
    再做列方向一维插值，最后在行方向插值。

    V2.2.2（C3）修复：行内空单元格（Excel 空白，读入为 None）不再让 None
    参与算术抛 TypeError（原被调用方 except 吞掉后静默返回 0）——
    取该行最近有效列值参与插值；整行无有效值时抛出明确 ValueError。
    """
    # clamp row
    rv = max(row_keys[0], min(row_keys[-1], row_val))

    # 找行区间
    ri = 0
    for i in range(len(row_keys) - 1):
        if row_keys[i] <= rv <= row_keys[i + 1]:
            ri = i
            break

    def _row_interp(r_idx, col):
        """对给定行做列方向一维插值（自动 clamp）"""
        r_key = row_keys[r_idx]
        max_c = row_col_max.get(r_key, col_keys[-1])
        c = min(col, max_c)  # 逐行限界
        vals = matrix[r_idx]

        def _valid_at(j):
            """j 位置的值；None（Excel 空白）时向两侧就近取该行最近有效值"""
            if 0 <= j < len(vals) and vals[j] is not None:
                return vals[j]
            for step in range(1, len(vals)):
                lo, hi = j - step, j + step
                if lo >= 0 and vals[lo] is not None:
                    return vals[lo]
                if hi < len(vals) and vals[hi] is not None:
                    return vals[hi]
            return None

        def _require(v, what):
            if v is None:
                raise ValueError(f'承载力表行 {r_key} 列 {what} 无有效值（整行空白或数据缺失）')
            return v

        if c <= col_keys[0]:
            return _require(_valid_at(0), col_keys[0])
        if c >= col_keys[-1]:
            return _require(_valid_at(len(col_keys) - 1), col_keys[-1])
        for j in range(len(col_keys) - 1):
            if col_keys[j] <= c <= col_keys[j + 1]:
                dx = col_keys[j + 1] - col_keys[j]
                if dx > 0:
                    v1 = _require(_valid_at(j), col_keys[j])
                    v2 = _require(_valid_at(j + 1), col_keys[j + 1])
                    return v1 - (c - col_keys[j]) * (v1 - v2) / dx
                return _require(_valid_at(j), col_keys[j])
        return _require(_valid_at(len(col_keys) - 1), col_keys[-1])

    v1 = _row_interp(ri, col_val)
    v2 = _row_interp(ri + 1, col_val)

    # 行间线性插值
    x1, x2 = row_keys[ri], row_keys[ri + 1]
    if x2 > x1:
        tx = (rv - x1) / (x2 - x1)
        return v1 + (v2 - v1) * tx
    return v1


def _compute_q4clay(e_avg, il_avg):
    """Q4黏性土：e × IL → σ0 — 数据实时读取 Excel

    V3.0.3（Agent2-3）：e=1.1 行数据自 IL=0.2 起（Excel 与 TB10012-2001
    表 D.0.1-5 原文同值），IL∈(0,0.2) 无格值。规范原文无法确认该行是 IL=0
    左对齐还是自 0.2 起——保守沿用 V2.2.2（C3）"取该行最近有效列值"策略
    （IL<0.2 → 160），与规范可能差值 ≤10kPa，属保守方向可接受。
    """
    tbl = _load_lookup_from_excel()
    if not tbl: return 0
    t = tbl['q4clay']
    return _ragged_2d(e_avg, max(il_avg, 0), t['e'], t['il'], t['matrix'], t['il_max'])


def _compute_silt(e_avg, w_avg):
    """粉土：e × W → σ0 — 数据实时读取 Excel"""
    tbl = _load_lookup_from_excel()
    if not tbl: return 0
    t = tbl['silt']
    return _ragged_2d(e_avg, max(w_avg, 10), t['e'], t['w'], t['matrix'], t['w_max'])


def _compute_soft(w_avg):
    """软土：W → σ0 — 数据实时读取 Excel"""
    tbl = _load_lookup_from_excel()
    if not tbl: return 0
    t = tbl['soft']
    return _interp_1d(max(w_avg, t['w'][0]), t['w'], t['sigma'])


def _compute_q3clay(es_avg, e_avg=None, il_avg=None):
    """Q3以前黏性土：Es → σ0 — 数据实时读取 Excel

    V3.0.3（Agent2-2）：Es<10MPa 回落 Q4 黏性土表（e×IL）。
    依据：TB10012-2001 表 D.0.1-6（《铁桥地规》表4.1.2-6）注 2"当压缩模量小于
    10MPa 时，其基本承载力可按黏性土表 D.0.1-5 确定"（知识库《岩土工程常用数据
    公式汇总》第十七章铁路专题同注；规范原文 PDF P83 文本核对）。
    此前 Es<10 一律 clamp 到 Es 表首行 380（如 Es=8 按规范途径应回落 Q4 表
    e=0.9/IL=0.5 → 210，原实现系统性高估）。
    实现：Es<10 且提供 e>0（IL 缺失按 0，与 Q4 分支 clamp 口径一致）时按 Q4 表
    _compute_q4clay 计算；e 缺失（无法确定回落值）时保留原 clamp 行为并记日志。
    """
    tbl = _load_lookup_from_excel()
    if not tbl: return 0
    if es_avg < 10.0:
        if e_avg is not None and e_avg > 0:
            return _compute_q4clay(e_avg, il_avg if il_avg is not None else 0)
        _log.warning('Q3以前黏性土 Es=%r<10MPa 需按 D.0.1-6 注2 回落 Q4 表，但 e 缺失，保留 Es 表首值 380', es_avg)
        t = tbl['q3clay']
        return t['sigma'][0]
    t = tbl['q3clay']
    return _interp_1d(min(es_avg, t['es'][-1]), t['es'], t['sigma'])


def _compute_residual(es_avg):
    """残积黏性土：Es → σ0 — 数据实时读取 Excel"""
    tbl = _load_lookup_from_excel()
    if not tbl: return 0
    t = tbl['residual']
    return _interp_1d(min(es_avg, t['es'][-1]), t['es'], t['sigma'])


# V3.0.3（Agent2-5）：表 D.0.1-9（《铁桥地规》表4.1.2-8）注 1"非饱和 Q3 新黄土，
# 当 0.85<e<0.95 时 σ0 值可提高 10%"——未实现：当前数据流无饱和度/地层时代字段，
# 无法判定"非饱和 Q3"（仅凭名称含 黄土 无法区分）；项目表（铁建承载力计算表.xlsx
# 新黄土 sheet）亦未含该提高条款。仅注释标注，不改行为。
def _compute_loess_new(wl_avg, e_avg, w_avg):
    """新黄土：WL × e × W — 数据实时读取 Excel

    V3.0.3（Agent2-1）：WL 分块间线性内插。
    依据：TB10012-2001 表 D.0.1-9（《铁桥地规》表4.1.2-8）注 3"表中括号内数值
    用于内插取值"（知识库《岩土工程常用数据公式汇总》第十七章铁路专题同注；
    注册岩土考试辅导亦按 WL 分块内插查表）。
    此前 WL∈(24,28)/(28,32) 直接取低分块（如 WL=30/e=0.9/W=20 取 WL=28 块 160，
    WL=32 块同格 200，差 40kPa），系统性偏低。
    实现：先对各 WL 分块（24/28/32）按 e/W 网格双线性插值得 v24/v28/v32
    （块内 e/W 插值与块间 WL 插值共存），再对 WL 线性内插；WL≤24/≥32 取端块，
    WL=24/28/32 恰值分别取对应分块值（插值式在端点自然退化为块值）。
    NaN WL 按 S10 口径返回 None（不再静默落入低分块）。
    """
    tbl = _load_lookup_from_excel()
    if not tbl: return 0
    t = tbl['loess']
    if _is_nan(wl_avg):
        _log.warning('_compute_loess_new 收到 NaN WL，返回 None（e=%r w=%r）', e_avg, w_avg)
        return None
    # WL 分块节点：与 _load_lookup_from_excel 的 m24/m28/m32 结构一一对应；
    # TOML「承载力计算.黄土液限分段」[28,32] 即块间边界（_loess_wl_thresholds），
    # 与 Excel 块结构不一致时以 Excel 为准（仅告警，不臆改行为）。
    wl_lo, wl_mid, wl_hi = 24, 28, 32
    if list(_loess_wl_thresholds) != [wl_mid, wl_hi]:
        _log.warning('黄土液限分段 TOML=%r 与 Excel 新黄土块 WL=24/28/32 不一致，按 Excel 块结构内插', _loess_wl_thresholds)
    v24 = _interp_2d(e_avg, w_avg, t['e'], t['w24'], t['m24'])
    v28 = _interp_2d(e_avg, w_avg, t['e'], t['w28'], t['m28'])
    v32 = _interp_2d(e_avg, w_avg, t['e'], t['w32'], t['m32'])
    if _is_nan(v24) or _is_nan(v28) or _is_nan(v32):
        return None
    if wl_avg <= wl_lo:
        return v24
    if wl_avg <= wl_mid:  # (24, 28]
        return v24 + (v28 - v24) * (wl_avg - wl_lo) / (wl_mid - wl_lo)
    if wl_avg <= wl_hi:   # (28, 32]
        return v28 + (v32 - v28) * (wl_avg - wl_mid) / (wl_hi - wl_mid)
    return v32


def _compute_loess_old(e_avg, w_over_wl):
    """老黄土：e × W/WL 离散查表 — 数据实时读取 Excel"""
    tbl = _load_lookup_from_excel()
    if not tbl: return 0
    m = tbl['loess_old']['matrix']
    if e_avg < _loess_old_e_thresholds[0]:     e_idx = 0
    elif e_avg <= _loess_old_e_thresholds[1]:  e_idx = 1
    elif e_avg <= _loess_old_e_thresholds[2]:  e_idx = 2
    else:                                      e_idx = 3
    if w_over_wl < 0.6:       r_idx = 0
    elif w_over_wl <= 0.8:    r_idx = 1
    else:                     r_idx = 2
    return m[r_idx][e_idx]


# ============================================================
# 岩土名称 → 计算函数 匹配
# ============================================================

# 按关键词匹配（顺序敏感，先匹配更具体的）
# V2.2.3（S3）：老黏土/Q3以前黏性土接线——'老黏土'/'老黏性土'/'Q3' 关键词必须位于
# 任何含 '黏土'/'黏性土' 子串的关键词之前（否则被子串吞掉），使该类地层走
# _compute_q3clay（按压缩模量 Es 查表插值）。
# 规范来源：TB 10012-2001 铁路工程地质勘察规范 P104（Es=10/15/20/25/30/35/40 →
# σ0=380/430/470/510/550/580/620），与「铁建承载力计算表.xlsx」Q3以前黏性土 sheet 一致。
SOIL_TYPE_PATTERNS = [
    ('淤泥', 'soft'),
    ('泥炭', 'soft'),
    ('软土', 'soft'),
    ('残积黏性土', 'residual'),
    ('残积粉质黏土', 'residual'),
    ('新黄土', 'loess_new'),
    ('老黄土', 'loess_old'),
    ('黄土', 'loess_new'),       # 默认新黄土
    ('粉土', 'silt'),
    ('砂质粉土', 'silt'),
    ('黏质粉土', 'silt'),
    ('老黏土', 'q3clay'),
    ('老黏性土', 'q3clay'),
    # V2.2.4（E1）：'Q3' 不再是裸关键词——仅当名称同时含 '黏土'/'黏性土' 字样
    # （复合条件）才归 q3clay；'Q3粉砂/Q3圆砾土/Q3碎石/Q3砂岩/人工填土(Q3)' 等
    # 非黏性土不再误判。位置保持在 粉土 族之后、'粉质黏土' 之前（'Q3黏质粉土' 仍走粉土）。
    ('Q3', 'q3clay'),
    ('粉质黏土', 'q4clay'),
    ('黏土', 'q4clay'),
    ('砂质黏性土', 'q4clay'),
    ('砾质黏性土', 'q4clay'),
    ('黏性土', 'q4clay'),
]


def classify_soil_for_bearing(name: str) -> str:
    """根据岩土名称返回承载力计算类型"""
    if not name:
        return ''
    for keyword, soil_type in SOIL_TYPE_PATTERNS:
        if keyword == 'Q3':
            # E1 复合条件：'Q3' 仅当名称同时含 黏土/黏性土/黏质 时才归 q3clay
            # （Q3粉质黏土/Q3eol粉质黏土/粉砂质黏土(Q3)/Q32粉质黏土 均命中；
            #  V2.2.5（E3）补 '黏质'：'Q3黏质土' 是黏性土同义简称，V2.2.4 起漏判
            #  (q3clay→'')；'Q3黏质粉土' 仍被更早的 粉土 关键词截走归 silt；
            #  Q3粉砂/Q3中砂/Q3圆砾土/Q3砂岩/人工填土(Q3) 落入后续关键词或返回 ''）
            if 'Q3' in name and ('黏土' in name or '黏性土' in name or '黏质' in name):
                return soil_type
            continue
        if keyword in name:
            return soil_type
    return ''


def _fmt_num(x):
    """数值格式化：整数去小数，小数最多保留 3 位并去尾零"""
    if x is None:
        return '0'
    try:
        xf = float(x)
    except (ValueError, TypeError):
        return str(x)
    if abs(xf - round(xf)) < 1e-9:
        return str(int(round(xf)))
    return f'{xf:.3f}'.rstrip('0').rstrip('.')


def _fmt_sigma(v):
    """承载力数值格式化：整数去小数，否则保留 1 位小数"""
    if abs(v - round(v)) < 1e-9:
        return str(int(round(v)))
    return f'{v:.1f}'


def compute_bearing_capacity_detail(name: str, indicators: dict):
    """返回 (σ0原始插值 float, 查表列展示文本 str)

    文本格式如 'e=0.951 IL=0.49 查表：σ0=191.1'，用于 Excel 查表列展示。
    不适用/无法计算返回 (0.0, '')。
    """
    stype = classify_soil_for_bearing(name)
    if not stype:
        return 0.0, ''

    # V2.2.3（S10）：入参指标统一数值化——None/非法值/NaN 按 0 处理
    # （NaN 不再流入插值层被静默 clamp 到边界；0 值由下方各分支按
    # "指标缺失 → 无法计算"返回空值，并保留异常日志路径）
    def _num(v):
        if v is None:
            return 0
        try:
            f = float(v)
        except (ValueError, TypeError):
            return 0
        return f if f == f else 0

    e = _num(indicators.get('e'))
    il = _num(indicators.get('IL'))
    w = _num(indicators.get('W'))
    es = _num(indicators.get('Es'))
    wl = _num(indicators.get('WL'))

    try:
        if stype == 'q4clay':
            if e <= 0:
                return 0.0, ''
            # IL<0（坚硬状态）按 IL=0 档查表（_compute_q4clay 内部 max(il_avg,0) clamp），
            # 不再整层拒绝导致 S 列恒为空
            val = _compute_q4clay(e, il)
            desc = f'e={_fmt_num(e)} IL={_fmt_num(il)}'
        elif stype == 'silt':
            if e <= 0 or w <= 0:
                return 0.0, ''
            val = _compute_silt(e, w)
            desc = f'e={_fmt_num(e)} W={_fmt_num(w)}'
        elif stype == 'soft':
            if w <= 0:
                return 0.0, ''
            val = _compute_soft(w)
            desc = f'W={_fmt_num(w)}'
        elif stype == 'residual':
            if es <= 0:
                return 0.0, ''
            val = _compute_residual(es)
            desc = f'Es={_fmt_num(es)}'
        elif stype == 'q3clay':
            if es <= 0:
                return 0.0, ''
            val = _compute_q3clay(es, e, il)  # V3.0.3（Agent2-2）传 e/IL 供 Es<10 回落 Q4 表
            desc = f'Es={_fmt_num(es)}'
        elif stype == 'loess_new':
            if e <= 0 or w <= 0 or wl <= 0:
                return 0.0, ''
            val = _compute_loess_new(wl, e, w)
            desc = f'WL={_fmt_num(wl)} e={_fmt_num(e)} W={_fmt_num(w)}'
        elif stype == 'loess_old':
            if e <= 0 or w <= 0 or wl <= 0:
                return 0.0, ''
            wo = (w / wl) if wl > 0 else 0
            val = _compute_loess_old(e, wo)
            desc = f'e={_fmt_num(e)} W/WL={_fmt_num(wo)}'
        else:
            return 0.0, ''
    except Exception as e:
        # V2.2.2（C3）：禁止静默吞错——查表/插值异常记日志后按"无法计算"返回 0，
        # 避免用户看到 S 列空白却无任何痕迹（原来 TypeError 被吞掉静默返回 0）。
        _log.exception('compute_bearing_capacity_detail(%r) 异常: %s', name, e)
        return 0.0, ''

    if val is None:  # V2.2.3（S10）：插值层检测到 NaN 返回 None → 按无法计算处理
        return 0.0, ''

    return float(val), f'{desc} 查表：σ0={_fmt_sigma(val)}'


def compute_bearing_capacity(name: str, indicators: dict) -> int:
    """根据岩土名称和指标统计均值计算基本承载力 σ0 (kPa)

    indicators: {'e': 孔隙比均值, 'IL': 液指均值, 'W': 含水率均值,
                 'Es': 压缩模量均值, 'WL': 液限均值}

    返回 int kPa；不适用返回 0
    """
    v, _ = compute_bearing_capacity_detail(name, indicators)
    return int(round(v)) if v else 0


# ============================================================
# 地层参数表 K 值匹配（列 T）— 与 S 列完全独立
# ============================================================



def _norm_param_no(v):
    """地层编号归一化（与 rule_engine._norm_state_key 口径一致）

    '01'/'1.0'/数值 1.0 → '1'；非整数（如 '1.5'）与非数值（如 'A'）保留原串。
    V2.2.3（S9）：字符串前导零/浮点写法同样归一化，避免参数表编号精确匹配漏配。
    """
    s = str(v).strip()
    try:
        f = float(s)
    except (ValueError, TypeError):
        return s
    return str(int(f)) if f == int(f) else s


def _load_param_table():
    """加载统一地层标准状态参数表，返回 [(编号, 子编号, 岩性, 状态, K), ...]

    读取 参数/地层标准状态参数表.xlsx
    列1=主层编号, 列2=亚层, 列5=岩土名称, 列12=承载力, 列13=原始状态描述
    """
    base = os.path.dirname(os.path.dirname(os.path.dirname(__file__)))
    path = os.path.join(base, '参数', '地层标准状态参数表.xlsx')

    if not os.path.exists(path):
        return []

    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    result = []
    last_lithology = ''

    for r in range(2, ws.max_row + 1):
        b = ws.cell(row=r, column=1).value   # 主层编号
        sub = ws.cell(row=r, column=2).value  # 亚层
        i = ws.cell(row=r, column=5).value   # 岩土名称
        h = ws.cell(row=r, column=13).value  # 原始状态描述
        k = ws.cell(row=r, column=12).value  # 承载力

        if b is None:
            continue

        # V2.2.3（S9）：编号归一化与 rule_engine._norm_state_key 口径一致
        # （'01'/'1.0'/数值 1.0 → '1'；非整数如 '1.5' 保留原串），
        # 消除前导零/浮点地层编号在 lookup_param_sigma 精确匹配时的静默漏配
        b_key = _norm_param_no(b)
        sub_key = _norm_param_no(sub) if sub is not None else ''

        if i is not None:
            last_lithology = str(i).strip()
        i_key = last_lithology if last_lithology else ''
        h_key = str(h).strip() if h else ''

        try:
            k_val = int(float(k))
        except (ValueError, TypeError):
            k_val = None

        if k_val is not None and k_val > 0:
            result.append((b_key, sub_key, i_key, h_key, k_val))

    wb.close()
    return result


def _norm(s):
    """名称归一化：去掉中英文括号（用于岩性比较，不计 ''()()）"""
    return s.replace('（', '').replace('）', '').replace('(', '').replace(')', '')


def _name_match_exact(search: str, candidate: str):
    """精确匹配：search 是否等于 candidate，或等于其顿号分隔后的某个独立成分

    '灰岩' 匹配 '白云质灰岩、灰质白云岩、灰岩、白云岩'（组件相等）
    '灰岩' 不匹配 '炭质灰岩'（非组件，无分隔符）
    """
    sn = _norm(search)
    cn = _norm(candidate) if candidate else ''
    if not sn or not cn:
        return False
    if sn == cn:
        return True
    # 复合岩性拆分（全角顿号、逗号）
    for part in re.split(r'[、，,]', cn):
        if sn == part.strip():
            return True
    return False


def lookup_param_sigma(layer_label: str, rock_name: str, plasticity: str) -> int:
    """从广清永标准地层参数表匹配 K 值（基本承载力建议值）

    优先按地层编号（编号+子编号）精确映射；编号无匹配时回退岩性+状态匹配。

    layer_label: 如 "8-12"（编号=8, 子编号=12）、"2-31"（编号=2, 子编号=31）
    """
    table = _load_param_table()
    if not table:
        return 0

    label = str(layer_label).strip() if layer_label else ''
    name = str(rock_name).strip() if rock_name else ''
    state = str(plasticity).strip() if plasticity else ''

    # 提取编号与子编号：'8-12' -> base='8', sub='12'
    parts = label.split('-')
    base_raw = parts[0].strip()
    sub_raw = parts[1].strip() if len(parts) > 1 else ''
    # V2.2.4（E2）：查询端与加载端 _load_param_table 共用 _norm_param_no 归一化，
    # '01-1'/'1.0-1' 与归一化后的表键 ('1','1') 精确匹配（V2.2.3 只归一化加载端、
    # 查询端用原始编号导致前导零/浮点编号静默漏配，T 列显示 /）
    base = _norm_param_no(base_raw) if base_raw else ''
    sub = _norm_param_no(sub_raw) if sub_raw else ''
    if not base:
        return 0

    # ---- 优先：按 (编号, 子编号) 精确映射 ----
    for k_label, k_sub, k_name, k_state, k_val in table:
        if k_label == base and k_sub == sub:
            return k_val

    # ---- 回退：按编号 + 岩性 + 状态匹配 ----
    name_norm = _norm(name)

    # 1) 编号 + 岩性精确匹配 + 状态相等
    for k_label, k_sub, k_name, k_state, k_val in table:
        if k_label != base:
            continue
        if _name_match_exact(name, k_name) and k_state == state:
            return k_val

    # 2) 评分遍历
    best_val = 0
    best_score = -1
    best_kn_len = 9999
    for k_label, k_sub, k_name, k_state, k_val in table:
        if k_label != base:
            continue
        kn = _norm(k_name) if k_name else ''
        name_exact = False
        name_component = False
        kn_from_state = False
        if name_norm and kn:
            if _name_match_exact(name, k_name):
                if name_norm == kn:
                    name_exact = True
                else:
                    name_component = True
            else:
                ks = _norm(k_state) if k_state else ''
                if name_norm and ks:
                    if _name_match_exact(name, k_state):
                        kn = ks
                        name_exact = (name_norm == ks)
                        name_component = not name_exact
                        kn_from_state = True
                    else:
                        continue
                else:
                    continue
        if state and k_state and not kn_from_state:
            if not (state == k_state or state in k_state or k_state in state):
                continue
        score = 0
        if name_norm and kn:
            if name_exact:
                score += 10
            elif name_component:
                score += 5
        if state and k_state:
            if state == k_state:
                score += 10
            elif state in k_state:
                score += 5
            elif k_state in state:
                score += 3
        if score > best_score or (score == best_score and len(kn) <= best_kn_len):
            best_score = score
            best_val = k_val
            best_kn_len = len(kn)

    return best_val
